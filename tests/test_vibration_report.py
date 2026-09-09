import io
import os
import shutil
import tempfile
import unittest
from unittest import mock

from src import domain_measurements as dm
from src import vibration_report_notes as notes_store
from src import vibration_standards as vs
from src import vibration_report as report


class IsoZoneTests(unittest.TestCase):
    def test_group_1_rigid_limits_match_the_plant_form(self):
        # FORM.JRG.F.05.001 prints A >=0, B >=2.3, C >=4.5, D >=7.1 for this class.
        self.assertEqual(vs.zone_limits("ISO 10816-3 (GROUP 1 RIGID)"), (2.3, 4.5, 7.1))

    def test_each_class_in_the_asset_register_resolves(self):
        for equipment_class in (
            "ISO 10816-3 (GROUP 1 RIGID)",
            "ISO 10816-3 (GROUP 2 RIGID)",
            "ISO 10816-3 (GROUP 2 FLEXIBLE)",
            "ISO 10816-2 (GROUP 2 RIGID)",
        ):
            limits = vs.zone_limits(equipment_class)
            self.assertEqual(len(limits), 3)
            self.assertLess(limits[0], limits[1])
            self.assertLess(limits[1], limits[2])

    def test_class_without_a_group_falls_back_to_the_strictest(self):
        # "ISO 10816-2" names no group; judging it leniently would be the
        # dangerous direction, so it must land on the strictest table.
        self.assertEqual(vs.canon_equipment_class("ISO 10816-2"), "GROUP 2 RIGID")
        self.assertEqual(vs.zone_limits("ISO 10816-2"), (1.4, 2.8, 4.5))

    def test_same_reading_is_judged_differently_per_class(self):
        # 1.78 mm/s - the worst point in the plant's own example report.
        group_1 = vs.evaluate_overall(1.78, "ISO 10816-3 (GROUP 1 RIGID)")
        group_2 = vs.evaluate_overall(1.78, "ISO 10816-3 (GROUP 2 RIGID)")
        self.assertEqual(group_1["zone"], "A")
        self.assertEqual(group_1["status"], "NORMAL")
        self.assertEqual(group_2["zone"], "B")

    def test_zone_boundaries_are_inclusive_upward(self):
        self.assertEqual(vs.evaluate_overall(7.1, "ISO 10816-3 (GROUP 1 RIGID)")["zone"], "D")
        self.assertEqual(vs.evaluate_overall(7.09, "ISO 10816-3 (GROUP 1 RIGID)")["zone"], "C")


class ShockPulseTests(unittest.TestCase):
    def test_delta_matches_the_plant_worked_example(self):
        # Page 2 of the form: max -24, carpet -31 -> delta 7.
        self.assertEqual(vs.evaluate_shock_pulse(-24, -31)["delta"], 7)
        self.assertEqual(vs.evaluate_shock_pulse(-29, -29)["delta"], 0)

    def test_thresholds_follow_the_form(self):
        self.assertEqual(vs.evaluate_shock_pulse(20, 12)["carpet_status"], "WARNING")
        self.assertEqual(vs.evaluate_shock_pulse(20, 16)["carpet_status"], "ALARM")
        self.assertEqual(vs.evaluate_shock_pulse(26, 5)["max_status"], "WARNING")
        self.assertEqual(vs.evaluate_shock_pulse(36, 5)["max_status"], "ALARM")

    def test_overall_status_takes_the_worse_of_the_two(self):
        self.assertEqual(vs.evaluate_shock_pulse(36, 2)["status"], "ALARM")
        self.assertEqual(vs.evaluate_shock_pulse(2, 12)["status"], "WARNING")
        self.assertEqual(vs.evaluate_shock_pulse(2, 2)["status"], "NORMAL")

    def test_lubrication_reading_follows_the_delta_rule(self):
        self.assertEqual(vs.evaluate_shock_pulse(-24, -31)["lubrication"], "Kelebihan greasing")
        self.assertEqual(vs.evaluate_shock_pulse(0, -15)["lubrication"], "Kekurangan greasing")
        self.assertEqual(vs.evaluate_shock_pulse(0, -10)["lubrication"], "Greasing ideal")

    def test_missing_reading_is_reported_not_guessed(self):
        result = vs.evaluate_shock_pulse(None, -31)
        self.assertIsNone(result["delta"])
        self.assertEqual(result["status"], "TIDAK ADA DATA")


class ReportTestCase(unittest.TestCase):
    """Seeds the exact readings printed in the plant's own example report so
    the generated output can be checked against a known-good document."""

    EQUIPMENT = "Motor ID Fan 1#1"
    MONTHS = {
        "2024-07-15": [0.28, 0.71, 0.55, 0.40, 0.90, 0.61, 0.45, 0.64, 1.69,
                       0.45, 0.64, 1.64, 0.36, 0.61, 0.59, 0.67, 1.69, 1.19],
        "2024-08-15": [0.28, 0.58, 0.57, 0.28, 0.77, 0.63, 0.35, 0.53, 1.96,
                       0.43, 0.60, 0.69, 0.38, 0.59, 0.63, 1.13, 2.24, 1.42],
        "2024-09-09": [0.26, 0.53, 0.47, 0.34, 0.63, 0.52, 0.49, 0.59, 1.78,
                       0.61, 0.59, 0.58, 0.59, 0.68, 0.48, 0.37, 0.72, 1.68],
    }
    SHOCK_PULSE = {"sp_max_1": -29, "sp_carpet_1": -29, "sp_max_2": -28, "sp_carpet_2": -29,
                   "sp_max_3": -24, "sp_carpet_3": -31, "sp_max_4": -23, "sp_carpet_4": -31}

    def setUp(self):
        self.root = tempfile.mkdtemp(prefix="pple_vibrep_")
        self._env = mock.patch.dict(os.environ, {"MCSA_DATA_DIR": self.root})
        self._env.start()

        rows = []
        for test_date, values in self.MONTHS.items():
            for key, value in zip(report.POINT_KEYS, values):
                rows.append({"equipment": self.EQUIPMENT, "unit_name": "UNIT 1",
                             "test_date": test_date, "parameter": key, "value": value, "uom": "mm/s"})
        for key, value in self.SHOCK_PULSE.items():
            rows.append({"equipment": self.EQUIPMENT, "test_date": "2024-09-09",
                         "parameter": key, "value": value, "uom": "dB"})
        dm.append_measurements("VIBRASI", rows)

    def tearDown(self):
        self._env.stop()
        shutil.rmtree(self.root, ignore_errors=True)


class MonthlyTableTests(ReportTestCase):
    def test_one_row_per_month_newest_last(self):
        table = report.monthly_overall_table(self.EQUIPMENT)
        self.assertEqual(list(table["BULAN"]), ["Juli-24", "Agustus-24", "September-24"])

    def test_values_match_the_source_report_with_indonesian_decimals(self):
        table = report.monthly_overall_table(self.EQUIPMENT)
        september = table[table["BULAN"] == "September-24"].iloc[0]
        self.assertEqual(september["3A"], "1,78")
        self.assertEqual(september["1V"], "0,26")

    def test_all_eighteen_points_are_columns(self):
        table = report.monthly_overall_table(self.EQUIPMENT)
        self.assertEqual(len(table.columns), 19)  # BULAN + 18 points
        self.assertIn("6A", table.columns)

    def test_months_are_limited_to_the_requested_window(self):
        self.assertEqual(len(report.monthly_overall_table(self.EQUIPMENT, months=2)), 2)

    def test_unknown_equipment_yields_an_empty_table_not_an_error(self):
        self.assertTrue(report.monthly_overall_table("Tidak Ada").empty)


class ShockPulseRowTests(ReportTestCase):
    def test_deltas_match_the_source_report(self):
        rows = report.shock_pulse_rows(report.latest_readings(self.EQUIPMENT))
        self.assertEqual([row["delta"] for row in rows], [0, 1, 7, 8])

    def test_bearings_without_readings_are_marked_no_data(self):
        rows = report.shock_pulse_rows({"sp_max_1": -29, "sp_carpet_1": -29})
        self.assertEqual(rows[0]["status"], "NORMAL")
        self.assertEqual(rows[1]["status"], "TIDAK ADA DATA")


class DraftNotesTests(ReportTestCase):
    def test_draft_names_the_worst_point_like_the_plant_does(self):
        draft = notes_store.draft_notes(
            self.EQUIPMENT, "ISO 10816-3 (GROUP 1 RIGID)",
            report.latest_readings(self.EQUIPMENT),
            report.shock_pulse_rows(report.latest_readings(self.EQUIPMENT)),
        )
        self.assertIn("bearing 3 arah axial", draft["keterangan"])
        self.assertIn("1,78 mm/s", draft["analisa"])
        self.assertIn("NORMAL", draft["analisa"])
        self.assertIn("- Lakukan monitoring vibrasi secara rutin dan berkala", draft["rekomendasi"])

    def test_healthy_shock_pulse_does_not_add_greasing_noise(self):
        # All four bearings read far below warning; the plant's own report
        # raises no greasing action for them, and neither should the draft.
        draft = notes_store.draft_notes(
            self.EQUIPMENT, "ISO 10816-3 (GROUP 1 RIGID)",
            report.latest_readings(self.EQUIPMENT),
            report.shock_pulse_rows(report.latest_readings(self.EQUIPMENT)),
        )
        self.assertNotIn("greasing", draft["rekomendasi"].lower())

    def test_elevated_shock_pulse_does_raise_a_greasing_check(self):
        shock_pulse = report.shock_pulse_rows({"sp_max_1": 40, "sp_carpet_1": 20})
        draft = notes_store.draft_notes(
            self.EQUIPMENT, "ISO 10816-3 (GROUP 1 RIGID)",
            report.latest_readings(self.EQUIPMENT), shock_pulse,
        )
        self.assertIn("pelumasan bearing 1", draft["rekomendasi"])

    def test_severe_vibration_escalates_the_recommendation(self):
        draft = notes_store.draft_notes(
            "X", "ISO 10816-3 (GROUP 1 RIGID)", {"pt1_v": 9.0}, [],
        )
        self.assertIn("SOP", draft["rekomendasi"])

    def test_no_readings_yields_empty_draft_rather_than_invented_text(self):
        draft = notes_store.draft_notes("X", "ISO 10816-3 (GROUP 1 RIGID)", {}, [])
        self.assertEqual(set(draft.values()), {""})


class NotesStoreTests(ReportTestCase):
    def test_saved_notes_survive_and_override_the_draft(self):
        notes_store.save_notes(self.EQUIPMENT, "2024-09-09", {
            "keterangan": "catatan lapangan", "analisa": "analisa manual", "rekomendasi": "rekomendasi manual",
        })
        context = report.build_report_context(self.EQUIPMENT)
        self.assertEqual(context["notes"]["analisa"], "analisa manual")

    def test_unknown_fields_are_not_persisted(self):
        saved = notes_store.save_notes(self.EQUIPMENT, "2024-09-09", {"analisa": "x", "rahasia": "y"})
        self.assertNotIn("rahasia", saved)

    def test_notes_are_keyed_per_equipment_and_date(self):
        notes_store.save_notes(self.EQUIPMENT, "2024-09-09", {"analisa": "september"})
        self.assertEqual(notes_store.get_notes(self.EQUIPMENT, "2024-08-15")["analisa"], "")


class DocxTests(ReportTestCase):
    def _text(self):
        from docx import Document

        document = Document(io.BytesIO(report.build_docx(self.EQUIPMENT)))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)

    def test_document_carries_the_form_letterhead_and_number(self):
        text = self._text()
        self.assertIn("PT. INDONESIA POWER", text)
        self.assertIn("LAPORAN PDM TEKNOLOGI VIBRASI", text)
        self.assertIn("DETAIL REPORT VIBRASI", text)
        self.assertIn(report.FORM_NUMBER, text)

    def test_document_contains_the_monthly_table_values(self):
        text = self._text()
        self.assertIn("September-24", text)
        self.assertIn("1,78", text)

    def test_document_contains_shock_pulse_and_narrative(self):
        text = self._text()
        self.assertIn("DATA SHOCK PULSE", text)
        self.assertIn("ANALISA", text)
        self.assertIn("REKOMENDASI", text)

    def test_document_states_what_is_not_included(self):
        # Photos and spectrum plots have no data source yet; the report must
        # say so rather than leave a silent gap.
        self.assertIn("belum tersedia di sistem", self._text())

    def test_identity_table_keeps_a_genuine_zero_instead_of_dashing_it_out(self):
        # `value or "-"` would treat a real 0 (e.g. PM WEEK 0, an unscheduled
        # asset) as falsy and print "-" - it must print "0" instead.
        with mock.patch.object(report, "find_asset", return_value={"pm_week": 0, "kks": "LJ10H"}):
            text = self._text()
        self.assertIn("0", text.split("PM WEEK", 1)[1].split("\n")[1])

    def test_document_builds_for_equipment_missing_from_the_asset_register(self):
        dm.append_measurements("VIBRASI", [
            {"equipment": "Mesin Tanpa Register", "test_date": "2024-09-09", "parameter": "pt1_v", "value": 1.0},
        ])
        payload = report.build_docx("Mesin Tanpa Register")
        self.assertGreater(len(payload), 100)


if __name__ == "__main__":
    unittest.main()
