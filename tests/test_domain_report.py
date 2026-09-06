import io
import os
import shutil
import tempfile
import unittest
from datetime import date
from unittest import mock

import pandas as pd

from src import domain_measurements as dm
from src import domain_report as report


class DomainReportTestCase(unittest.TestCase):
    def setUp(self):
        self.root = tempfile.mkdtemp(prefix="pple_report_")
        self._env = mock.patch.dict(os.environ, {"MCSA_DATA_DIR": self.root})
        self._env.start()
        dm.append_measurements("VIBRASI", [
            {"equipment": "CWP 1A", "unit_name": "UNIT 1", "test_date": "2026-03-05", "parameter": "overall_rms", "value": 2.1, "uom": "mm/s"},
            {"equipment": "CWP 1A", "unit_name": "UNIT 1", "test_date": "2026-04-05", "parameter": "overall_rms", "value": 8.4, "uom": "mm/s"},
            {"equipment": "CWP 1B", "unit_name": "UNIT 1", "test_date": "2026-04-05", "parameter": "overall_rms", "value": 1.5, "uom": "mm/s"},
        ])
        dm.append_measurements("DGA", [
            {"equipment": "TRF-1", "test_date": "2026-04-10", "parameter": "h2", "value": 40, "uom": "ppm"},
            {"equipment": "TRF-1", "test_date": "2026-04-10", "parameter": "c2h2", "value": 3.0, "uom": "ppm"},
        ])

    def tearDown(self):
        self._env.stop()
        shutil.rmtree(self.root, ignore_errors=True)


class CollectTests(DomainReportTestCase):
    def test_collect_scopes_to_the_period(self):
        data = report.collect_domain_data("VIBRASI", date(2026, 4, 1), date(2026, 4, 30))
        self.assertEqual(data["reading_count"], 2)
        self.assertEqual(data["equipment_count"], 2)

    def test_collect_runs_the_real_agent_per_equipment(self):
        data = report.collect_domain_data("VIBRASI", date(2026, 1, 1), date(2026, 12, 31))
        verdicts = {item["equipment"]: item for item in data["verdicts"]}
        self.assertEqual(set(verdicts), {"CWP 1A", "CWP 1B"})
        # 8.4 mm/s is far beyond ISO 10816-3 zone C; 1.5 mm/s is healthy. The
        # verdicts must differ, proving stored values drove the evaluation.
        self.assertNotEqual(verdicts["CWP 1A"]["condition"], verdicts["CWP 1B"]["condition"])

    def test_empty_period_reports_zero_rather_than_being_omitted(self):
        data = report.collect_domain_data("VIBRASI", date(2020, 1, 1), date(2020, 12, 31))
        self.assertEqual(data["reading_count"], 0)
        self.assertEqual(data["verdicts"], [])
        self.assertEqual(data["label"], "Vibrasi")


class CsvReportTests(DomainReportTestCase):
    def test_csv_contains_the_domain_column_and_rows(self):
        raw = report.build_csv(["VIBRASI"], date(2026, 1, 1), date(2026, 12, 31))
        frame = pd.read_csv(io.BytesIO(raw))
        self.assertEqual(len(frame), 3)
        self.assertEqual(set(frame["domain"]), {"VIBRASI"})

    def test_csv_combines_multiple_domains(self):
        raw = report.build_csv(["VIBRASI", "DGA"], date(2026, 1, 1), date(2026, 12, 31))
        frame = pd.read_csv(io.BytesIO(raw))
        self.assertEqual(set(frame["domain"]), {"VIBRASI", "DGA"})
        self.assertEqual(len(frame), 5)

    def test_csv_for_empty_period_still_has_headers(self):
        raw = report.build_csv(["VIBRASI"], date(2020, 1, 1), date(2020, 2, 1))
        frame = pd.read_csv(io.BytesIO(raw))
        self.assertTrue(frame.empty)
        self.assertIn("domain", frame.columns)


class DocxReportTests(DomainReportTestCase):
    def test_docx_is_a_readable_document_with_domain_sections(self):
        from docx import Document

        raw = report.build_docx(["VIBRASI", "DGA"], date(2026, 1, 1), date(2026, 12, 31))
        document = Document(io.BytesIO(raw))
        text = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("Laporan Condition Monitoring", text)
        headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
        self.assertIn("Vibrasi", headings)
        self.assertIn("DGA", headings)

    def test_docx_states_empty_period_explicitly(self):
        from docx import Document

        raw = report.build_docx(["VIBRASI"], date(2020, 1, 1), date(2020, 2, 1))
        document = Document(io.BytesIO(raw))
        text = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("Tidak ada pengukuran pada periode ini.", text)

    def test_docx_detail_table_lists_every_reading(self):
        from docx import Document

        raw = report.build_docx(["VIBRASI"], date(2026, 1, 1), date(2026, 12, 31))
        document = Document(io.BytesIO(raw))
        detail = document.tables[-1]
        self.assertEqual(len(detail.rows), 1 + 3)  # header + 3 readings


class PptxReportTests(DomainReportTestCase):
    def test_pptx_has_a_title_slide_plus_one_per_domain(self):
        from pptx import Presentation

        raw = report.build_pptx(["VIBRASI", "DGA"], date(2026, 1, 1), date(2026, 12, 31))
        deck = Presentation(io.BytesIO(raw))
        self.assertEqual(len(deck.slides), 3)
        titles = [slide.shapes.title.text for slide in deck.slides]
        self.assertIn("Vibrasi", titles)
        self.assertIn("DGA", titles)

    def test_pptx_states_empty_period_on_the_slide(self):
        from pptx import Presentation

        raw = report.build_pptx(["VIBRASI"], date(2020, 1, 1), date(2020, 2, 1))
        deck = Presentation(io.BytesIO(raw))
        texts = [
            shape.text_frame.text
            for slide in deck.slides for shape in slide.shapes
            if shape.has_text_frame
        ]
        self.assertTrue(any("Tidak ada pengukuran" in text for text in texts))


class BundleTests(DomainReportTestCase):
    def test_bundle_returns_all_three_formats(self):
        bundle = report.build_report_bundle(["VIBRASI"], date(2026, 1, 1), date(2026, 12, 31))
        self.assertEqual(set(bundle), {"docx", "pptx", "csv"})
        for payload in bundle.values():
            self.assertIsInstance(payload, bytes)
            self.assertGreater(len(payload), 100)

    def test_bundle_can_be_narrowed_to_one_format(self):
        bundle = report.build_report_bundle(["VIBRASI"], date(2026, 1, 1), date(2026, 12, 31), formats=["csv"])
        self.assertEqual(set(bundle), {"csv"})

    def test_bundle_rejects_unknown_domain(self):
        with self.assertRaises(dm.UnknownDomainError):
            report.build_report_bundle(["NOPE"], date(2026, 1, 1), date(2026, 12, 31))


if __name__ == "__main__":
    unittest.main()
