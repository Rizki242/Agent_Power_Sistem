"""Tests for automated reports, meeting PPTX, and document date resolution."""

import unittest
from datetime import date
from io import BytesIO

from fastapi.testclient import TestClient

from api_server import app
from pple.api.routers.agents import detect_document_date
from src import domain_report as report


class AutomatedReportsTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)

    def test_detect_document_date_with_explicit_date(self):
        # Indonesian text date format
        text_id = "Laporan Pengujian Vibrasi Tanggal 15 September 2026 pada Unit 1 BFP 1A"
        dt, src, lbl = detect_document_date(text_id)
        self.assertEqual(dt, "2026-09-15")
        self.assertEqual(src, "document")
        self.assertIn("15", lbl)

        # ISO format
        text_iso = "Hasil uji laboratorium oli tanggal 2026-08-20 menunjukkan viskositas normal."
        dt_iso, src_iso, _ = detect_document_date(text_iso)
        self.assertEqual(dt_iso, "2026-08-20")
        self.assertEqual(src_iso, "document")

    def test_detect_document_date_without_date_falls_back_to_upload_date(self):
        text_no_date = "Pengujian vibrasi rutin pada motor CWP tanpa informasi tanggal."
        dt, src, lbl = detect_document_date(text_no_date)
        self.assertEqual(src, "upload_date")
        self.assertIn("Data merujuk pada tanggal upload", lbl)
        self.assertEqual(dt, date.today().strftime("%Y-%m-%d"))

    def test_get_automated_reports_summary_endpoint(self):
        resp = self.client.get("/api/reports/automated/summary")
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertIn("PLTU Jeranjang", data["plant"])
        self.assertEqual(data["current_sample_month"], "September 2026")
        self.assertTrue(len(data["sample_automated_reports"]) >= 4)
        self.assertTrue(len(data["modules_status_september_2026"]) == 6)

        # Ensure standby badge is generated when a module has no data
        for mod in data["modules_status_september_2026"]:
            if not mod["has_data"]:
                self.assertIn("STANDBY", mod["status_label"])
                self.assertIsNotNone(mod["standby_badge"])

    def test_download_weekly_report_endpoint(self):
        resp = self.client.get("/api/reports/automated/download/weekly/MCSA?format=docx&week=3")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("application/vnd.openxmlformats-officedocument", resp.headers["content-type"])
        self.assertGreater(len(resp.content), 500)

    def test_download_meeting_pptx_endpoint(self):
        resp = self.client.get("/api/reports/automated/download/meeting-pptx?year=2026&month=9")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("presentationml.presentation", resp.headers["content-type"])
        self.assertGreater(len(resp.content), 1000)

    def test_build_meeting_pptx_structure(self):
        from pptx import Presentation

        raw = report.build_meeting_pptx(
            ["MCSA", "VIBRASI", "DGA", "PD", "TRIBOLOGY", "THERMAL"],
            start=date(2026, 9, 1),
            end=date(2026, 9, 30),
        )
        deck = Presentation(BytesIO(raw))
        # Cover + Matrix + 6 domain slides + Action items = 9 slides
        self.assertGreaterEqual(len(deck.slides), 8)

        # Check for presence of standby warning text in the deck
        all_texts = [
            shape.text_frame.text
            for slide in deck.slides
            for shape in slide.shapes
            if shape.has_text_frame
        ]
        self.assertTrue(any("BELUM ADA DATA PENGUJIAN / STANDBY" in t for t in all_texts))

    def test_download_monthly_per_module_endpoints(self):
        modules = ["MCSA", "VIBRASI", "DGA", "PD", "TRIBOLOGY", "THERMAL"]
        for mod in modules:
            # Test docx
            resp_docx = self.client.get(f"/api/reports/automated/download/monthly/{mod}?format=docx&year=2026&month=9")
            self.assertEqual(resp_docx.status_code, 200)
            self.assertIn("application/vnd.openxmlformats-officedocument.wordprocessingml.document", resp_docx.headers["content-type"])
            self.assertGreater(len(resp_docx.content), 500)
            self.assertIn(mod, resp_docx.headers.get("content-disposition", ""))

            # Test pptx
            resp_pptx = self.client.get(f"/api/reports/automated/download/monthly/{mod}?format=pptx&year=2026&month=9")
            self.assertEqual(resp_pptx.status_code, 200)
            self.assertIn("application/vnd.openxmlformats-officedocument.presentationml.presentation", resp_pptx.headers["content-type"])
            self.assertGreater(len(resp_pptx.content), 500)
            self.assertIn(mod, resp_pptx.headers.get("content-disposition", ""))

    def test_download_monthly_consolidated_reports(self):
        # Test all domains consolidated DOCX
        resp_docx = self.client.get("/api/reports/automated/download/monthly?format=docx&year=2026&month=9")
        self.assertEqual(resp_docx.status_code, 200)
        self.assertIn("wordprocessingml", resp_docx.headers["content-type"])
        self.assertGreater(len(resp_docx.content), 1000)

        # Test all domains consolidated PPTX
        resp_pptx = self.client.get("/api/reports/automated/download/monthly?format=pptx&year=2026&month=9")
        self.assertEqual(resp_pptx.status_code, 200)
        self.assertIn("presentationml", resp_pptx.headers["content-type"])
        self.assertGreater(len(resp_pptx.content), 1000)

        # Test all domains consolidated CSV
        resp_csv = self.client.get("/api/reports/automated/download/monthly?format=csv&year=2026&month=9")
        self.assertEqual(resp_csv.status_code, 200)
        self.assertIn("text/csv", resp_csv.headers["content-type"])
        self.assertGreater(len(resp_csv.content), 50)

