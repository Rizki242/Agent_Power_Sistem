"""Unit tests for Technology Examination (TE) FORM.JRG.F.05.006 Generator."""

import io
import unittest
from docx import Document

from src.te_generator import build_te_report_data, export_te_docx


class TestTechnologyExamination(unittest.TestCase):
    def test_build_te_report_data_structure(self):
        eq = "BOILER FEEDWATER PUMP 2 UNIT 3"
        data = build_te_report_data(
            equipment=eq,
            technology="Vibrasi dan Tribology",
            finding="Vibrasi tinggi dan NAS pelumas tinggi",
            status="Kuning",
        )

        self.assertEqual(data["equipment"], eq.upper())
        self.assertEqual(data["form_code"], "FORM.JRG.F.05.006")
        self.assertEqual(data["revision"], "01")
        self.assertIn("TE/CBM/UJPJRJ", data["doc_number"])
        self.assertEqual(data["technology"], "Vibrasi dan Tribology")
        self.assertEqual(data["status"], "Kuning")

        # Specifications check
        specs = data.get("specifications", {})
        self.assertIn("motor", specs)
        self.assertIn("driven", specs)
        self.assertIn("type_mfg", specs["motor"])
        self.assertIn("speed", specs["motor"])
        self.assertIn("power", specs["motor"])
        self.assertIn("inboard_bearing", specs["motor"])

        # Vibration data check
        vib = data.get("vibration_data", {})
        self.assertIn("points", vib)
        for pt in ("1V", "1H", "1A", "2V", "2H", "2A", "4V", "4A"):
            self.assertIn(pt, vib["points"])
        self.assertGreater(len(vib.get("spectrum_notes", [])), 0)

        # Thermal & Tribology check
        self.assertIn("thermal_data", data)
        self.assertIn("temperatures", data["thermal_data"])
        self.assertIn("tribology_data", data)
        self.assertIn("reference_oil", data["tribology_data"])
        self.assertIn("total_fe_ppm", data["tribology_data"])

        # Analysis, Conclusion, Recommendation, Sign-off
        self.assertIn("analysis", data)
        self.assertGreater(len(data.get("conclusions", [])), 0)
        self.assertGreater(len(data.get("recommendations", [])), 0)

        sign = data.get("sign_off", {})
        self.assertEqual(sign.get("acknowledged_title"), "SPS RSO")
        self.assertEqual(sign.get("prepared_title"), "Pelaksana PdM")

    def test_export_te_docx_generation(self):
        data = build_te_report_data(
            equipment="CWP 1 UNIT 1",
            technology="MCSA dan Vibrasi",
            finding="Indikasi Broken Rotor Bar Sideband -36dB",
            status="Merah",
        )
        docx_buffer = export_te_docx(data)
        self.assertIsInstance(docx_buffer, io.BytesIO)
        self.assertGreater(docx_buffer.getbuffer().nbytes, 10000)

        # Re-parse docx to confirm valid document structure
        doc = Document(docx_buffer)
        self.assertGreater(len(doc.paragraphs), 10)
        self.assertGreater(len(doc.tables), 5)

        # Confirm official IMS header exists in table 0
        t0_text = doc.tables[0].cell(0, 0).text
        self.assertIn("PT. INDONESIA POWER", t0_text)
        self.assertIn("PLTU JERANJANG", t0_text)
        self.assertIn("INTEGRATED MANAGEMENT SYSTEM", t0_text)


if __name__ == "__main__":
    unittest.main()

