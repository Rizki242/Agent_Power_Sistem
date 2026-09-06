"""DETAIL REPORT VIBRASI - the plant's own per-equipment report form.

Reproduces FORM.JRG.F.05.001 as a Word document: the identity and machine
specification block, the monthly 1V-6A overall table with the ISO 10816-3
limit row for that machine's own class, the shock-pulse table, and the
engineer's KETERANGAN / ANALISA / REKOMENDASI.

Where each part comes from - nothing here is invented:
  - identity + MOTOR/driven specs: the existing vibration asset register
    (src.vibration_data.load_vibration_assets, backed by the SQLite that was
    itself extracted from this same report set)
  - monthly 1V-6A readings and shock pulse: the canonical measurement store
    (src.domain_measurements), i.e. whatever was uploaded or typed in
  - zone limits and shock-pulse verdicts: src.vibration_standards
  - narrative: src.vibration_report_notes (agent draft, engineer-editable)

Deliberately out of scope for now: the equipment photograph, the measuring
point diagram, and the spectrum / time-signal plots on page 2. Those need
image and raw-waveform sources the app does not have yet; a section is left
marked as such rather than filled with a placeholder that could be mistaken
for real data.
"""

from __future__ import annotations

from datetime import date
from io import BytesIO
from typing import Any, Optional

import pandas as pd

from src import domain_measurements as dm
from src import vibration_report_notes as notes_store
from src.vibration_standards import evaluate_overall, evaluate_shock_pulse, zone_limits

POINT_KEYS = [f"pt{point}_{axis}" for point in range(1, 7) for axis in ("v", "h", "a")]
POINT_HEADERS = [f"{point}{axis.upper()}" for point in range(1, 7) for axis in ("v", "h", "a")]
SHOCK_PULSE_BEARINGS = (1, 2, 3, 4)

FORM_NUMBER = "FORM.JRG.F.05.001"
FORM_ISSUED = "1 Sep 2016"
FORM_REVISION = "00"


def find_asset(equipment: str) -> Optional[dict[str, Any]]:
    """The asset register row for this equipment, matched on name or KKS.

    Returns None when the register is missing or unreadable rather than
    raising: the specification block is a nice-to-have on the report, and a
    deployment pointed at a data volume without the vibration SQLite should
    still be able to print the measurements it does have.
    """
    from src.vibration_data import load_vibration_assets

    try:
        assets = load_vibration_assets()
    except Exception:
        return None
    if assets is None or assets.empty:
        return None

    wanted = str(equipment or "").strip().casefold()
    for column in ("equipment", "kks", "asset_id"):
        if column not in assets.columns:
            continue
        hit = assets[assets[column].astype(str).str.strip().str.casefold() == wanted]
        if not hit.empty:
            return hit.iloc[0].to_dict()
    return None


def monthly_overall_table(equipment: str, months: int = 3, end: Optional[date] = None) -> pd.DataFrame:
    """The 1V-6A table, one row per month, newest last.

    Mirrors the form's own three-month view. A month with no reading is
    omitted rather than shown as zero - a blank month means "not measured",
    which is not the same as "measured zero".
    """
    frame = dm.filter_measurements("VIBRASI", parameters=POINT_KEYS, equipment=equipment)
    if frame.empty:
        return pd.DataFrame(columns=["BULAN", *POINT_HEADERS])

    frame = frame[frame["test_date"].notna()].copy()
    if end is not None:
        frame = frame[frame["test_date"].dt.date <= end]
    if frame.empty:
        return pd.DataFrame(columns=["BULAN", *POINT_HEADERS])

    frame["_month"] = frame["test_date"].dt.to_period("M")
    latest_months = sorted(frame["_month"].unique())[-months:]
    frame = frame[frame["_month"].isin(latest_months)]

    rows = []
    for period in latest_months:
        scoped = frame[frame["_month"] == period]
        # Newest reading wins when a month holds more than one test.
        newest = scoped.sort_values("test_date").drop_duplicates(subset=["parameter"], keep="last")
        by_parameter = dict(zip(newest["parameter"], newest["value"]))
        row = {"BULAN": period.strftime("%B-%y")}
        for key, header in zip(POINT_KEYS, POINT_HEADERS):
            value = by_parameter.get(key)
            row[header] = "" if value is None or pd.isna(value) else f"{float(value):.2f}".replace(".", ",")
        rows.append(row)

    return pd.DataFrame(rows, columns=["BULAN", *POINT_HEADERS])


def latest_readings(equipment: str, end: Optional[date] = None) -> dict[str, float]:
    """Newest numeric value of every stored parameter for this equipment."""
    frame = dm.filter_measurements("VIBRASI", equipment=equipment)
    if frame.empty:
        return {}
    frame = frame[frame["test_date"].notna()]
    if end is not None:
        frame = frame[frame["test_date"].dt.date <= end]
    if frame.empty:
        return {}

    newest = frame.sort_values("test_date").drop_duplicates(subset=["parameter"], keep="last")
    readings: dict[str, float] = {}
    for _, row in newest.iterrows():
        value = row.get("value")
        if pd.notna(value):
            readings[str(row["parameter"])] = float(value)
    return readings


def shock_pulse_rows(readings: dict[str, float]) -> list[dict[str, Any]]:
    """Per-bearing shock-pulse verdicts, one entry per bearing 1-4."""
    rows = []
    for bearing in SHOCK_PULSE_BEARINGS:
        verdict = evaluate_shock_pulse(
            readings.get(f"sp_max_{bearing}"),
            readings.get(f"sp_carpet_{bearing}"),
        )
        rows.append({"bearing": bearing, **verdict})
    return rows


def build_report_context(equipment: str, end: Optional[date] = None) -> dict[str, Any]:
    """Everything one printed report needs, gathered but not yet laid out."""
    asset = find_asset(equipment) or {}
    equipment_class = asset.get("equipment_class_normalized") or asset.get("equipment_class_source") or ""
    readings = latest_readings(equipment, end=end)
    shock_pulse = shock_pulse_rows(readings)

    point_values = {key: value for key, value in readings.items() if key.startswith("pt")}
    verdict = None
    if point_values:
        worst = max(point_values.values())
        verdict = evaluate_overall(worst, equipment_class)

    saved = notes_store.get_notes(equipment, _measurement_date(equipment, end))
    if not any(saved.values()):
        saved = notes_store.draft_notes(equipment, equipment_class, readings, shock_pulse)

    return {
        "equipment": equipment,
        "asset": asset,
        "equipment_class": equipment_class,
        "readings": readings,
        "shock_pulse": shock_pulse,
        "verdict": verdict,
        "monthly": monthly_overall_table(equipment, end=end),
        "notes": saved,
        "measurement_date": _measurement_date(equipment, end),
    }


def _measurement_date(equipment: str, end: Optional[date] = None):
    frame = dm.filter_measurements("VIBRASI", equipment=equipment)
    if frame.empty or frame["test_date"].isna().all():
        return None
    frame = frame[frame["test_date"].notna()]
    if end is not None:
        frame = frame[frame["test_date"].dt.date <= end]
    if frame.empty:
        return None
    return frame["test_date"].max().date()


def build_docx(equipment: str, end: Optional[date] = None) -> bytes:
    """Render the DETAIL REPORT VIBRASI form for one equipment."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    context = build_report_context(equipment, end=end)
    asset = context["asset"]
    document = Document()

    # --- letterhead ------------------------------------------------------
    header = document.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run("PT. INDONESIA POWER\nPLTU JERANJANG OMU\nBIDANG ENJINERING\nLAPORAN PDM TEKNOLOGI VIBRASI")
    run.bold = True
    run.font.size = Pt(12)

    control = document.add_paragraph(
        f"No. Dok : {FORM_NUMBER}    Tgl Terbit : {FORM_ISSUED}    REV : {FORM_REVISION}"
    )
    control.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in control.runs:
        run.font.size = Pt(8)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("DETAIL REPORT VIBRASI")
    run.bold = True
    run.font.size = Pt(14)

    # --- identity --------------------------------------------------------
    measurement_date = context["measurement_date"]
    identity = [
        ("EQUIPMENT", context["equipment"]),
        ("KKS", asset.get("kks", "-")),
        ("EQUIPMENT CLASS", context["equipment_class"] or "-"),
        ("PM WEEK", asset.get("pm_week", "-")),
        ("TGL PENGUKURAN", "-" if measurement_date is None else f"{measurement_date:%d-%b-%y}"),
        ("STATUS VIBRASI", (context["verdict"] or {}).get("status", asset.get("status_vibrasi", "-"))),
    ]
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in identity:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = str(value or "-")

    # --- machine specification (MOTOR | driven) --------------------------
    document.add_paragraph()
    spec_heading = document.add_paragraph()
    spec_heading.add_run("SPESIFIKASI").bold = True

    component_1 = str(asset.get("component_1") or "MOTOR")
    component_2 = str(asset.get("component_2") or "DRIVEN")
    spec_rows = [
        ("Type, Mfg", asset.get("c1_type_mfg"), asset.get("c2_type_mfg")),
        ("Speed", asset.get("c1_speed"), asset.get("c2_speed")),
        ("Power / Capacity", asset.get("c1_power"), asset.get("c2_capacity") or asset.get("c2_power")),
        ("Bearing Type", asset.get("c1_bearing_type"), asset.get("c2_bearing_type")),
        ("Inboard Bearing", asset.get("c1_inboard_bearing"), asset.get("c2_inboard_bearing")),
        ("Outboard Bearing", asset.get("c1_outboard_bearing"), asset.get("c2_onboard_bearing")),
        ("Rotor Bar / Total Blade", asset.get("c1_rotor_bar"), asset.get("c2_total_blade")),
        ("Foundation", asset.get("c1_foundation"), ""),
    ]
    spec_table = document.add_table(rows=1, cols=3)
    spec_table.style = "Table Grid"
    for cell, text in zip(spec_table.rows[0].cells, ("", component_1, component_2)):
        cell.text = text
    for label, left, right in spec_rows:
        cells = spec_table.add_row().cells
        cells[0].text = label
        cells[1].text = "-" if left in (None, "") or pd.isna(left) else str(left)
        cells[2].text = "-" if right in (None, "") or (isinstance(right, float) and pd.isna(right)) else str(right)

    # --- monthly overall table ------------------------------------------
    document.add_paragraph()
    overall_heading = document.add_paragraph()
    overall_heading.add_run("DATA VIBRASI OVERALL").bold = True

    monthly = context["monthly"]
    if monthly.empty:
        document.add_paragraph("Tidak ada pengukuran tersimpan untuk equipment ini.")
    else:
        overall_table = document.add_table(rows=1, cols=len(monthly.columns))
        overall_table.style = "Table Grid"
        for cell, column in zip(overall_table.rows[0].cells, monthly.columns):
            cell.text = str(column)
        for _, row in monthly.iterrows():
            cells = overall_table.add_row().cells
            for cell, column in zip(cells, monthly.columns):
                cell.text = str(row[column])
        for row in overall_table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(7)

    # --- ISO limits for this machine's own class -------------------------
    a_b, b_c, c_d = zone_limits(context["equipment_class"])
    document.add_paragraph()
    limit_heading = document.add_paragraph()
    limit_heading.add_run(f"STANDART {context['equipment_class'] or 'ISO 10816-3'} - ALARM LIMIT (mm/s)").bold = True
    limit_table = document.add_table(rows=2, cols=4)
    limit_table.style = "Table Grid"
    for cell, text in zip(limit_table.rows[0].cells, ("A", "B", "C", "D")):
        cell.text = text
    for cell, text in zip(limit_table.rows[1].cells, (">=0", f">= {a_b}", f">= {b_c}", f">= {c_d}")):
        cell.text = text

    document.add_paragraph()
    document.add_paragraph("KETERANGAN :")
    document.add_paragraph(context["notes"].get("keterangan", "") or "-")

    # --- page 2: shock pulse + narrative ---------------------------------
    document.add_page_break()
    sp_heading = document.add_paragraph()
    sp_heading.add_run("DATA SHOCK PULSE").bold = True

    sp_table = document.add_table(rows=1, cols=5)
    sp_table.style = "Table Grid"
    for cell, text in zip(sp_table.rows[0].cells, ("BEARING", "1", "2", "3", "4")):
        cell.text = text
    for label, key in (("Max (dB)", "max_db"), ("Carpet (dB)", "carpet_db"), ("Delta Shock Pulse", "delta")):
        cells = sp_table.add_row().cells
        cells[0].text = label
        for index, entry in enumerate(context["shock_pulse"], start=1):
            value = entry.get(key)
            cells[index].text = "-" if value is None else f"{value:g}"
    status_cells = sp_table.add_row().cells
    status_cells[0].text = "Status"
    for index, entry in enumerate(context["shock_pulse"], start=1):
        status_cells[index].text = str(entry.get("status", "-"))

    document.add_paragraph(
        "KETERANGAN : Carpet warning >=10 dB, alarm >=15 dB. Max warning >=25 dB, alarm >=35 dB. "
        "Best practice delta shock pulse = 10 (<10 kelebihan greasing, >10 kekurangan greasing)."
    )

    document.add_paragraph()
    analysis_heading = document.add_paragraph()
    analysis_heading.add_run("ANALISA").bold = True
    document.add_paragraph(context["notes"].get("analisa", "") or "-")

    recommendation_heading = document.add_paragraph()
    recommendation_heading.add_run("REKOMENDASI").bold = True
    document.add_paragraph(context["notes"].get("rekomendasi", "") or "-")

    document.add_paragraph()
    pending = document.add_paragraph(
        "Foto equipment, diagram titik pengukuran, serta grafik spectrum dan time signal "
        "belum tersedia di sistem dan tidak disertakan pada laporan ini."
    )
    for run in pending.runs:
        run.italic = True
        run.font.size = Pt(8)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
