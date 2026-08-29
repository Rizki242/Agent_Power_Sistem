"""Generator laporan PowerPoint MCSA.

Struktur deck (16:9):
  1  Cover
  2  Ringkasan Eksekutif   (KPI card + donut + sorotan)
  3  Top 10 Risiko         (tabel + bar skor)
  4  Distribusi per Unit   (stacked column)
  5+ Divider per Unit/Voltage, lalu slide detail tiap equipment
     - Normal/Standby : tabel parameter full-width
     - Alarm/High     : tabel + panel gambar spektrum + slide tren
  N  Rekomendasi & Tindak Lanjut
  N+1 Metodologi & Catatan Pengambilan Data

Seluruh warna, font, dan geometri berasal dari `src/ppt_theme.py`.
"""

import io
import json
import os
from datetime import datetime

import pandas as pd
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_DATA_LABEL_POSITION, XL_TICK_MARK

from src import ppt_theme as T
from src.analytics import build_risk_summary
from src.ppt_assets import SpectrumImageIndex

_EQUIPMENT_MAPPING_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'equipment_mapping.json')

# Parameter yang ditampilkan lebih dulu pada tabel detail.
PREFERRED_ORDER = [
    'Kondisi', 'Bearing', 'Load', 'power factor', 'Real Power',
    'Voltage 1', 'Voltage 2', 'Voltage 3',
    'Current 1', 'Current 2', 'Current 3',
    'Dev Voltage', 'Dev Current',
    'THD Voltage %', 'THD Current %',
    'Rotorbar', 'Rotorbar Health', 'Rotorbar Level %', 'Rotorbar Severity Level',
    'Upper Sideband', 'Lower Sideband', 'Se Fund', 'Se Harm',
]

# Parameter yang di-plot pada slide tren.
TREND_PARAMS = ['THD Voltage %', 'THD Current %', 'Dev Current', 'Upper Sideband']

METODOLOGI_LINES = [
    "Pengambilan data memakai ATPOLL II; pastikan koneksi clamp arus/tegangan aman dan polaritas benar.",
    "Ambil data pada beban memadai (>=40-60% nameplate) agar analisa rotor bar reliabel.",
    "Rekam minimal 30-60 detik untuk stabilitas spektrum; hindari transien start/stop.",
    "Catat metadata: Unit, Voltage, beban, tanggal/jam, dan kondisi operasi.",
    "Konsistenkan titik ukur dan durasi rekaman agar hasil antar sesi dapat dibandingkan.",
]

TINDAK_LANJUT_LINES = [
    "Normal - lanjutkan operasi, trending berkala, simpan rekaman sebagai baseline.",
    "Alarm - inspeksi terarah (koneksi fasa, terminal, ground), survei vibrasi, naikkan frekuensi trending.",
    "High - tindakan segera: survei vibrasi menyeluruh, inspeksi koneksi/terminal, evaluasi beban dan supply.",
]


# --------------------------------------------------------------------------
# Helper data
# --------------------------------------------------------------------------

def _load_equipment_mapping():
    try:
        with open(_EQUIPMENT_MAPPING_PATH, 'r', encoding='utf-8') as fh:
            return json.load(fh) or {}
    except Exception:
        return {}


def _equipment_long_name(equipment, mapping):
    """Terjemahkan prefix kode equipment ke nama panjang (BFP -> Boiler Feed Pump)."""
    code = ''.join(ch for ch in str(equipment or '') if ch.isalpha()).upper()
    if not code:
        return None
    for key in sorted(mapping, key=len, reverse=True):
        if code.startswith(key.upper()):
            return mapping[key]
    return None


def _unit_limit(param_name: str):
    p = (param_name or '').strip().lower()
    if p in {'dev voltage', 'dev voltage %'}:
        return '%', 'Alarm >1%, High >2%'
    if p in {'dev current', 'dev current %'}:
        return '%', 'Alarm >5%, High >10%'
    if p in {'thd voltage %', 'thd current %'}:
        return '%', 'Alarm >5%, High >8%'
    if p == 'load':
        return '%', 'Valid >=40%, Monitoring >=20%'
    if p.startswith('current'):
        return 'A', ''
    if p.startswith('voltage'):
        return 'V', ''
    if p in {'real power', 'power'}:
        return 'kW', ''
    if p in {'upper sideband', 'lower sideband'}:
        return 'dB', 'Alarm > -54 dB, High > -45 dB'
    if p in {'rotorbar level %'}:
        return '%', ''
    return '', ''


def _format_value(param_name: str, raw_value, num_value):
    v = num_value if (isinstance(num_value, (int, float)) and not pd.isna(num_value)) else raw_value
    if isinstance(v, (int, float)) and not pd.isna(v):
        p = (param_name or '').strip().lower()
        if p in {'dev voltage', 'dev voltage %', 'dev current', 'dev current %',
                 'thd voltage %', 'thd current %', 'load', 'rotorbar level %',
                 'upper sideband', 'lower sideband'}:
            return f"{float(v):.1f}"
        if p.startswith('current') or p.startswith('voltage'):
            return f"{float(v):.2f}"
        if p in {'real power', 'power'}:
            return f"{float(v):.3f}"
        return f"{v}"
    if isinstance(v, str):
        s = ' '.join(v.replace('\r', ' ').replace('\n', ' ').split())
        return s if s else "-"
    return str(v) if v is not None else "-"


def _clean_text(value):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ''
    s = str(value).replace('\r', ' ').replace('\n', ' ')
    # Beberapa sel CSV menyimpan "\n" sebagai dua karakter literal.
    s = s.replace('\\n', ' ').replace('\\r', ' ')
    s = ' '.join(s.split())
    return '' if s.lower() in {'nan', 'none'} else s


def _shorten(text, limit):
    """Potong teks agar sel tabel tidak membungkus dan menaikkan tinggi baris."""
    s = str(text or '')
    return s if len(s) <= limit else s[:max(1, limit - 1)].rstrip() + '…'


def _thresholds():
    """Ambang Alarm/High per parameter, dibaca dari config bila tersedia."""
    try:
        from src.standards import load_thresholds_config
        cfg = load_thresholds_config() or {}
    except Exception:
        cfg = {}

    def g(section, key, default):
        try:
            return float((cfg.get(section) or {}).get(key, default))
        except (TypeError, ValueError):
            return float(default)

    return {
        'THD Voltage %': (g('thd_voltage', 'alarm_pct', 5), g('thd_voltage', 'high_pct', 8), 'higher'),
        'THD Current %': (g('thd_current', 'alarm_pct', 5), g('thd_current', 'high_pct', 8), 'higher'),
        'Dev Voltage': (g('unbalance_voltage', 'alarm_pct', 1), g('unbalance_voltage', 'high_pct', 2), 'higher'),
        'Dev Current': (g('unbalance_current', 'alarm_pct', 5), g('unbalance_current', 'high_pct', 10), 'higher'),
        # Ambang sideband mengikuti panduan MCSA/ESA di README (bagus < -54 dB).
        'Upper Sideband': (-54.0, -45.0, 'higher'),
        'Lower Sideband': (-54.0, -45.0, 'higher'),
    }


def _param_status(param_name, value, thresholds):
    """Status satu parameter numerik terhadap ambangnya, atau None bila N/A."""
    spec = thresholds.get(param_name)
    if spec is None:
        return None
    try:
        v = float(value)
    except (TypeError, ValueError):
        return None
    if pd.isna(v):
        return None
    alarm, high, direction = spec
    if direction == 'higher':
        if v > high:
            return 'High'
        if v > alarm:
            return 'Alarm'
        return 'Normal'
    if v < high:
        return 'High'
    if v < alarm:
        return 'Alarm'
    return 'Normal'


def _numeric(row):
    v = row.get('Value')
    if v is None or (isinstance(v, float) and pd.isna(v)):
        v = row.get('Raw_Value')
    try:
        f = float(v)
        return None if pd.isna(f) else f
    except (TypeError, ValueError):
        return None


# --------------------------------------------------------------------------
# Deck helper
# --------------------------------------------------------------------------

class _Deck:
    """Bungkus Presentation: nomor halaman dan footer otomatis."""

    def __init__(self, footer_text=''):
        self.prs = T.new_presentation()
        self.footer = footer_text
        self.page = 0

    def cover(self, title, subtitle_lines=None, footer=None):
        self.page += 1
        return T.add_cover_slide(self.prs, title, subtitle_lines, footer)

    def slide(self, title, meta_right=None):
        self.page += 1
        s = T.add_slide(self.prs, title, meta_right=meta_right, footer_left=self.footer)
        T.add_page_number(s, self.page)
        return s

    def divider(self, title, subtitle=None, lines=None):
        self.page += 1
        return T.add_divider_slide(self.prs, title, subtitle, lines)

    def save(self):
        out = io.BytesIO()
        self.prs.save(out)
        out.seek(0)
        return out


# --------------------------------------------------------------------------
# Slide builders
# --------------------------------------------------------------------------

def _build_summary_slide(deck, status_counts, total_eq, risk_rows, meta_right,
                         eq_infos=None):
    slide = deck.slide("Ringkasan Eksekutif", meta_right=meta_right)

    cards = ['Normal', 'Alarm', 'High', 'Standby']
    gap = Inches(0.25)
    card_w = int((T.CONTENT_W - 3 * gap) / 4)
    card_h = Inches(1.15)
    for i, st in enumerate(cards):
        T.kpi_card(slide, T.MARGIN_X + i * (card_w + gap), T.CONTENT_TOP,
                   card_w, card_h, int(status_counts.get(st, 0)), st, st)

    row2_top = T.CONTENT_TOP + card_h + Inches(0.25)
    row2_h = T.CONTENT_BOTTOM - row2_top
    chart_w = Inches(5.6)

    labels = [s for s in T.STATUS_ORDER if int(status_counts.get(s, 0)) > 0]
    if labels:
        T.rounded_panel(slide, T.MARGIN_X, row2_top, chart_w, row2_h, fill=T.PANEL, line=T.BORDER)
        T.add_textbox(slide, T.MARGIN_X + Inches(0.25), row2_top + Inches(0.18),
                      chart_w - Inches(0.5), Inches(0.28),
                      "Distribusi Status Equipment", size=Pt(12), color=T.INK, bold=True)

        cd = CategoryChartData()
        cd.categories = labels
        cd.add_series('Equipment', [int(status_counts.get(s, 0)) for s in labels])
        gf = slide.shapes.add_chart(
            XL_CHART_TYPE.DOUGHNUT,
            T.MARGIN_X + Inches(0.15), row2_top + Inches(0.52),
            chart_w - Inches(0.30), row2_h - Inches(0.72), cd)
        chart = gf.chart
        T.style_chart(chart, legend=True, legend_pos=XL_LEGEND_POSITION.BOTTOM)
        plot = chart.plots[0]
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.show_percentage = True
        dl.show_value = False
        dl.number_format = '0%'
        dl.number_format_is_linked = False
        dl.font.size = Pt(9)
        dl.font.bold = True
        dl.font.color.rgb = T.WHITE
        series = chart.series[0]
        for idx, st in enumerate(labels):
            pt = series.points[idx]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = T.status_color(st)
            pt.format.line.color.rgb = T.WHITE
            pt.format.line.width = Pt(1.5)

    # Panel sorotan
    panel_x = T.MARGIN_X + chart_w + Inches(0.25)
    panel_w = T.CONTENT_W - chart_w - Inches(0.25)
    T.rounded_panel(slide, panel_x, row2_top, panel_w, row2_h, fill=T.PANEL, line=T.BORDER)
    T.add_textbox(slide, panel_x + Inches(0.3), row2_top + Inches(0.22),
                  panel_w - Inches(0.6), Inches(0.28),
                  "Sorotan", size=Pt(12), color=T.INK, bold=True)

    issues = int(status_counts.get('Alarm', 0)) + int(status_counts.get('High', 0))
    pct = (issues / total_eq * 100.0) if total_eq else 0.0
    highlights = [
        f"Total equipment dianalisa: {total_eq}",
        f"Equipment perlu perhatian (Alarm/High): {issues} ({pct:.0f}%)",
    ]
    if risk_rows:
        worst = risk_rows[0]
        highlights.append(
            f"Skor kesehatan terendah: {worst.get('Full_Name') or worst.get('Equipment')} "
            f"({worst.get('Health Score')})"
        )
        drivers = []
        for r in risk_rows:
            for d in str(r.get('Drivers') or '').split(','):
                d = d.strip()
                if d and d not in drivers:
                    drivers.append(d)
        if drivers:
            highlights.append("Pemicu dominan: " + ", ".join(drivers[:4]))
    if int(status_counts.get('Standby', 0)):
        highlights.append(
            f"{int(status_counts['Standby'])} equipment standby - data tidak terbarui periode ini."
        )
    highlights.append("Rincian tindak lanjut ada di slide Rekomendasi.")

    T.bullet_list(slide, panel_x + Inches(0.3), row2_top + Inches(0.62),
                  panel_w - Inches(0.6), Inches(1.55),
                  highlights, size=Pt(11.5))

    # Daftar equipment yang perlu perhatian, mengisi sisa panel.
    list_top = row2_top + Inches(2.25)
    T.rect(slide, panel_x + Inches(0.3), list_top - Inches(0.14),
           panel_w - Inches(0.6), Pt(1), fill=T.BORDER)
    T.add_textbox(slide, panel_x + Inches(0.3), list_top,
                  panel_w - Inches(0.6), Inches(0.26),
                  "Equipment Perlu Perhatian", size=Pt(11.5), color=T.INK, bold=True)

    attention = []
    seen_names = set()
    for st in ('High', 'Alarm'):
        for eq, info in sorted((eq_infos or {}).items(), key=lambda kv: str(kv[0])):
            if info['status'] != st:
                continue
            key = info['title_short'].strip().lower()
            if key in seen_names:
                continue
            seen_names.add(key)
            attention.append((st, info['title_short']))

    chip_top = list_top + Inches(0.36)
    chip_h = Inches(0.28)
    chip_gap = Inches(0.08)
    col_w = int((panel_w - Inches(0.7)) / 3)
    max_rows_chip = 3
    shown = attention[:max_rows_chip * 3]
    for i, (st, name) in enumerate(shown):
        col, row = i % 3, i // 3
        x = panel_x + Inches(0.3) + col * (col_w + Inches(0.05))
        y = chip_top + row * (chip_h + chip_gap)
        T.rect(slide, x, y, Inches(0.05), chip_h, fill=T.status_color(st))
        T.add_textbox(slide, x + Inches(0.13), y + Inches(0.04),
                      col_w - Inches(0.2), Inches(0.22),
                      _shorten(name, 22), size=Pt(9.5), color=T.INK)
    if len(attention) > len(shown):
        T.add_textbox(slide, panel_x + Inches(0.3),
                      chip_top + max_rows_chip * (chip_h + chip_gap),
                      panel_w - Inches(0.6), Inches(0.24),
                      f"+ {len(attention) - len(shown)} equipment lainnya",
                      size=Pt(9.5), color=T.MUTED)
    if not attention:
        T.add_textbox(slide, panel_x + Inches(0.3), chip_top,
                      panel_w - Inches(0.6), Inches(0.26),
                      "Tidak ada equipment berstatus Alarm atau High.",
                      size=Pt(10.5), color=T.MUTED)
    return slide


def _build_risk_slide(deck, risk_rows, meta_right):
    slide = deck.slide("Top 10 Equipment Risiko Tertinggi", meta_right=meta_right)
    if not risk_rows:
        T.add_textbox(slide, T.MARGIN_X, T.CONTENT_TOP, T.CONTENT_W, Inches(0.4),
                      "Data tidak cukup untuk menghitung skor risiko.",
                      size=Pt(12), color=T.MUTED)
        return slide

    widths = [Inches(3.30), Inches(2.05), Inches(0.95), Inches(2.30), Inches(3.63)]
    row_h = Inches(0.40)
    table = T.add_table(slide, len(risk_rows) + 1, 5, T.MARGIN_X, T.CONTENT_TOP,
                        T.CONTENT_W, row_h * (len(risk_rows) + 1))
    headers = ["Equipment", "Unit / Voltage", "Skor", "", "Pemicu"]
    for c, h in enumerate(headers):
        table.cell(0, c).text = h

    for i, row in enumerate(risk_rows, start=1):
        table.cell(i, 0).text = str(row.get('Full_Name') or row.get('Equipment') or '-')
        unit = _clean_text(row.get('Unit_Name'))
        volt = _clean_text(row.get('Voltage_Level'))
        table.cell(i, 1).text = " / ".join([p for p in [unit, volt] if p]) or '-'
        table.cell(i, 2).text = str(row.get('Health Score', '-'))
        table.cell(i, 3).text = ""
        table.cell(i, 4).text = _clean_text(row.get('Drivers')) or '-'

    T.style_table(table, col_widths=widths, font_size=Pt(10.5), row_h=row_h,
                  align=[PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.CENTER,
                         PP_ALIGN.LEFT, PP_ALIGN.LEFT])

    # Bar skor digambar di atas kolom kosong ke-4.
    bar_x = T.MARGIN_X + sum(int(w) for w in widths[:4]) - int(widths[3]) + Inches(0.10)
    bar_max = int(widths[3]) - Inches(0.20)
    for i, row in enumerate(risk_rows, start=1):
        try:
            score = max(0.0, min(100.0, float(row.get('Health Score', 0))))
        except (TypeError, ValueError):
            score = 0.0
        st = 'High' if score < 60 else ('Alarm' if score < 80 else 'Normal')
        top = T.CONTENT_TOP + row_h * i + Inches(0.12)
        T.rect(slide, bar_x, top, bar_max, Inches(0.16), fill=T.BORDER)
        if score > 0:
            T.rect(slide, bar_x, top, int(bar_max * score / 100.0), Inches(0.16),
                   fill=T.status_color(st))

    T.add_textbox(slide, T.MARGIN_X, T.CONTENT_TOP + row_h * (len(risk_rows) + 1) + Inches(0.18),
                  T.CONTENT_W, Inches(0.5),
                  "Skor 100 = tanpa temuan. Pengurangan berasal dari status Rotorbar, THD, "
                  "unbalance tegangan/arus, kondisi bearing, dan validitas beban.",
                  size=Pt(9.5), color=T.MUTED)
    return slide


def _build_distribution_slide(deck, cond_df, meta_right):
    groups = []
    for (unit, volt), grp in cond_df.groupby(['_unit', '_volt'], dropna=False):
        label = " / ".join([p for p in [str(unit), str(volt)] if p and p != 'nan']) or 'Lainnya'
        counts = grp['_status'].value_counts().to_dict()
        groups.append((label, counts))
    if not groups:
        return None
    groups.sort(key=lambda g: g[0])

    slide = deck.slide("Distribusi Status per Unit dan Level Tegangan", meta_right=meta_right)
    cd = CategoryChartData()
    cd.categories = [g[0] for g in groups]
    series_status = [s for s in T.STATUS_ORDER
                     if any(int(g[1].get(s, 0)) > 0 for g in groups)]
    for st in series_status:
        cd.add_series(st, [int(g[1].get(st, 0)) for g in groups])

    gf = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_STACKED,
        T.MARGIN_X, T.CONTENT_TOP, T.CONTENT_W, T.CONTENT_H - Inches(0.1), cd)
    chart = gf.chart
    T.style_chart(chart, legend=True, legend_pos=XL_LEGEND_POSITION.BOTTOM)
    for idx, st in enumerate(series_status):
        ser = chart.series[idx]
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = T.status_color(st)
        ser.format.line.fill.background()
    plot = chart.plots[0]
    plot.gap_width = 80
    plot.has_data_labels = True
    plot.data_labels.font.size = Pt(9)
    plot.data_labels.font.color.rgb = T.WHITE
    plot.data_labels.font.bold = True
    plot.data_labels.position = XL_DATA_LABEL_POSITION.CENTER
    # Format 4-seksi dengan seksi "zero" kosong: label bernilai 0 disembunyikan.
    plot.data_labels.number_format = '0;-0;;'
    plot.data_labels.number_format_is_linked = False
    try:
        chart.value_axis.has_major_gridlines = True
        chart.value_axis.major_gridlines.format.line.color.rgb = T.BORDER
        chart.value_axis.major_tick_mark = XL_TICK_MARK.NONE
        chart.category_axis.major_tick_mark = XL_TICK_MARK.NONE
    except Exception:
        pass
    return slide


def _param_rows(eq_data, thresholds):
    """Kembalikan list dict siap render + status per baris."""
    rows = []
    for _, row in eq_data.iterrows():
        p_name = str(row.get('Parameter', ''))
        num_val = row.get('Value') if 'Value' in row.index else None
        raw_val = row.get('Raw_Value') if 'Raw_Value' in row.index else None
        value_txt = _format_value(p_name, raw_val, num_val)

        unit_txt = _clean_text(row.get('Unit')) if 'Unit' in row.index else ''
        limit_txt = _clean_text(row.get('Limit')) if 'Limit' in row.index else ''
        if not unit_txt or not limit_txt:
            u_def, l_def = _unit_limit(p_name)
            unit_txt = unit_txt or u_def
            limit_txt = limit_txt or l_def

        status = _param_status(p_name, _numeric(row), thresholds)
        if status is None and p_name in {'Kondisi', 'Bearing', 'Rotorbar'}:
            st = T.canon_status(raw_val)
            status = st if st in {'Alarm', 'High'} else None

        rows.append({
            'param': p_name,
            'value': value_txt,
            'unit': unit_txt or '-',
            'limit': limit_txt or '-',
            'status': status if status in {'Alarm', 'High'} else None,
        })
    return rows


def _render_param_table(slide, rows, left, top, width, row_h, font_size,
                        compact=False):
    cols = 3 if compact else 4
    table = T.add_table(slide, len(rows) + 1, cols, left, top, width,
                        row_h * (len(rows) + 1))
    if compact:
        headers = ["Parameter", "Nilai", "Limit"]
        widths = [int(width * 0.42), int(width * 0.22), int(width * 0.36)]
        align = [PP_ALIGN.LEFT, PP_ALIGN.RIGHT, PP_ALIGN.LEFT]
    else:
        headers = ["Parameter", "Nilai Terakhir", "Satuan", "Limit"]
        widths = [int(width * 0.33), int(width * 0.19), int(width * 0.11), int(width * 0.37)]
        align = [PP_ALIGN.LEFT, PP_ALIGN.RIGHT, PP_ALIGN.CENTER, PP_ALIGN.LEFT]

    for c, h in enumerate(headers):
        table.cell(0, c).text = h

    highlight = {}
    for i, r in enumerate(rows, start=1):
        if compact:
            val = r['value']
            if r['unit'] and r['unit'] != '-':
                val = f"{val} {r['unit']}"
            table.cell(i, 0).text = _shorten(r['param'], 26)
            table.cell(i, 1).text = _shorten(val, 16)
            table.cell(i, 2).text = _shorten(r['limit'], 28)
        else:
            table.cell(i, 0).text = _shorten(r['param'], 44)
            table.cell(i, 1).text = _shorten(r['value'], 30)
            table.cell(i, 2).text = r['unit']
            table.cell(i, 3).text = _shorten(r['limit'], 58)
        if r['status']:
            highlight[i] = r['status']

    T.style_table(table, col_widths=widths, font_size=font_size, row_h=row_h,
                  align=align, highlight_rows=highlight)
    return table


def _render_spectrum_panel(slide, img_index, equipment, date, left, top, width, height):
    """Panel gambar spektrum bertumpuk. Return True bila ada gambar."""
    picks = img_index.find(equipment, limit=3, date=date)
    if not picks:
        return False

    T.rounded_panel(slide, left, top, width, height, fill=T.PANEL, line=T.BORDER)
    label_date = picks[0][1]
    T.add_textbox(slide, left + Inches(0.18), top + Inches(0.14),
                  width - Inches(0.36), Inches(0.24),
                  f"Spektrum - {label_date}", size=Pt(10.5), color=T.INK, bold=True)

    inner_top = top + Inches(0.46)
    inner_h = height - Inches(0.62)
    gap = Inches(0.10)
    n = len(picks)
    cell_h = int((inner_h - gap * (n - 1)) / n)
    max_w = width - Inches(0.36)

    for i, (name, _d) in enumerate(picks):
        stream = img_index.load(name)
        if stream is None:
            continue
        ar = img_index.aspect_ratio(name)
        w = int(min(max_w, cell_h * ar))
        h = int(w / ar) if ar else cell_h
        if h > cell_h:
            h = cell_h
            w = int(h * ar)
        x = left + int((width - w) / 2)
        y = inner_top + i * (cell_h + gap) + int((cell_h - h) / 2)
        try:
            slide.shapes.add_picture(stream, x, y, width=w, height=h)
        except Exception:
            continue
    return True


def _build_detail_slides(deck, equipment, eq_data, info, thresholds, img_index,
                         include_images, meta_right):
    rows = _param_rows(eq_data, thresholds)
    if not rows:
        return

    status = info['status']
    critical = status in {'Alarm', 'High'}
    has_images = bool(include_images and critical and img_index and img_index.has(equipment))

    title = info['title']
    sub_bits = [b for b in [info['unit'], info['volt'], info['date_txt']] if b]
    subtitle = "  ·  ".join(sub_bits)

    row_h = Inches(0.27)
    # Slide pertama menyisakan ruang untuk panel spektrum; slide lanjutan
    # memakai lebar penuh supaya tidak ada setengah slide kosong.
    first_rows, rest_rows = (17, 18) if has_images else (18, 18)

    chunks = [rows[:first_rows]] if rows else [[]]
    remainder = rows[first_rows:]
    while remainder:
        chunks.append(remainder[:rest_rows])
        remainder = remainder[rest_rows:]
    total = len(chunks)

    for part, chunk in enumerate(chunks, start=1):
        with_images = has_images and part == 1
        table_w = Inches(6.35) if with_images else T.CONTENT_W
        font_size = Pt(9.5) if with_images else Pt(10.5)

        suffix = f" ({part}/{total})" if total > 1 else ""
        slide = deck.slide(f"{title}{suffix}", meta_right=meta_right)
        T.status_chip(slide, T.SLIDE_W - T.MARGIN_X - Inches(1.35), Inches(0.46),
                      status, status)
        if subtitle:
            T.add_textbox(slide, T.MARGIN_X + Inches(0.20), Inches(0.96),
                          T.CONTENT_W - Inches(3.0), Inches(0.26),
                          subtitle, size=Pt(10.5), color=T.MUTED)

        _render_param_table(slide, chunk, T.MARGIN_X, T.CONTENT_TOP,
                            table_w, row_h, font_size, compact=with_images)

        if with_images:
            panel_x = T.MARGIN_X + table_w + Inches(0.25)
            panel_w = T.CONTENT_W - table_w - Inches(0.25)
            _render_spectrum_panel(slide, img_index, equipment, info['date'],
                                   panel_x, T.CONTENT_TOP, panel_w, T.CONTENT_H)


def _series_for(history_df, equipment, param, max_points=12):
    if history_df is None or history_df.empty:
        return [], []
    sub = history_df[(history_df['Equipment'].astype(str) == str(equipment)) &
                     (history_df['Parameter'].astype(str) == param)]
    if sub.empty:
        return [], []
    sub = sub.copy()
    sub['_d'] = pd.to_datetime(sub.get('Date'), errors='coerce')
    sub['_v'] = pd.to_numeric(sub.get('Value'), errors='coerce')
    sub['_v'] = sub['_v'].fillna(pd.to_numeric(sub.get('Raw_Value'), errors='coerce'))
    sub = sub.dropna(subset=['_d', '_v']).sort_values('_d')
    sub = sub.drop_duplicates(subset=['_d'], keep='last').tail(max_points)
    if len(sub) < 3:
        return [], []
    return [d.strftime('%d %b %y') for d in sub['_d']], [float(v) for v in sub['_v']]


def _add_trend_chart(slide, left, top, width, height, param, cats, vals, thresholds):
    T.rounded_panel(slide, left, top, width, height, fill=T.PANEL, line=T.BORDER)
    T.add_textbox(slide, left + Inches(0.18), top + Inches(0.12),
                  width - Inches(0.36), Inches(0.24),
                  param, size=Pt(11), color=T.INK, bold=True)

    if not cats:
        T.add_textbox(slide, left + Inches(0.18), top + height / 2 - Inches(0.15),
                      width - Inches(0.36), Inches(0.3),
                      "Data riwayat belum cukup (< 3 titik).",
                      size=Pt(10), color=T.MUTED, align=PP_ALIGN.CENTER)
        return

    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series(param, vals)
    spec = thresholds.get(param)
    if spec:
        alarm, high, _dir = spec
        cd.add_series('Batas Alarm', [alarm] * len(cats))
        cd.add_series('Batas High', [high] * len(cats))

    gf = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        left + Inches(0.10), top + Inches(0.40),
        width - Inches(0.20), height - Inches(0.52), cd)
    chart = gf.chart
    T.style_chart(chart, legend=True, legend_pos=XL_LEGEND_POSITION.BOTTOM, font_size=Pt(8))

    main = chart.series[0]
    main.format.line.color.rgb = T.BRAND
    main.format.line.width = Pt(2.25)
    main.smooth = False
    if len(chart.series) > 1:
        for ser, color in ((chart.series[1], T.status_color('Alarm')),
                           (chart.series[2], T.status_color('High'))):
            ser.format.line.color.rgb = color
            ser.format.line.width = Pt(1.25)
            ser.marker.style = -4142  # xlMarkerStyleNone
            ser.smooth = False
    try:
        chart.value_axis.has_major_gridlines = True
        chart.value_axis.major_gridlines.format.line.color.rgb = T.BORDER
        chart.value_axis.major_tick_mark = XL_TICK_MARK.NONE
        chart.category_axis.major_tick_mark = XL_TICK_MARK.NONE
    except Exception:
        pass


def _build_trend_slide(deck, equipment, info, history_df, thresholds, meta_right):
    series = {p: _series_for(history_df, equipment, p) for p in TREND_PARAMS}
    if not any(cats for cats, _ in series.values()):
        return None

    slide = deck.slide(f"Tren Parameter - {info['title_short']}", meta_right=meta_right)
    T.status_chip(slide, T.SLIDE_W - T.MARGIN_X - Inches(1.35), Inches(0.46),
                  info['status'], info['status'])

    gap = Inches(0.22)
    w = int((T.CONTENT_W - gap) / 2)
    h = int((T.CONTENT_H - gap) / 2)
    positions = [
        (T.MARGIN_X, T.CONTENT_TOP),
        (T.MARGIN_X + w + gap, T.CONTENT_TOP),
        (T.MARGIN_X, T.CONTENT_TOP + h + gap),
        (T.MARGIN_X + w + gap, T.CONTENT_TOP + h + gap),
    ]
    for (x, y), param in zip(positions, TREND_PARAMS):
        cats, vals = series[param]
        _add_trend_chart(slide, x, y, w, h, param, cats, vals, thresholds)
    return slide


def _build_recommendation_slides(deck, df_latest, eq_infos, meta_right):
    try:
        from src.standards import generate_esa_mcsa_quick_recommendations
    except Exception:
        generate_esa_mcsa_quick_recommendations = None

    buckets = {'High': [], 'Alarm': []}
    references = []
    if generate_esa_mcsa_quick_recommendations is not None:
        for eq, info in eq_infos.items():
            if info['status'] not in {'Alarm', 'High'}:
                continue
            eq_data = df_latest[df_latest['Equipment'] == eq]
            try:
                res = generate_esa_mcsa_quick_recommendations(eq_data) or {}
            except Exception:
                continue
            bucket = 'High' if info['status'] == 'High' else 'Alarm'
            for rec in (res.get('recommendations') or []):
                rec = _clean_text(rec)
                if rec and rec not in buckets[bucket]:
                    buckets[bucket].append(rec)
            for ref in (res.get('references') or []):
                ref = _clean_text(ref)
                if ref and ref not in references:
                    references.append(ref)

    slide = deck.slide("Rekomendasi dan Tindak Lanjut", meta_right=meta_right)
    col_gap = Inches(0.25)
    col_w = int((T.CONTENT_W - col_gap) / 2)

    def _column(x, heading, status, items, empty_text):
        T.rounded_panel(slide, x, T.CONTENT_TOP, col_w, T.CONTENT_H,
                        fill=T.PANEL, line=T.BORDER)
        T.rect(slide, x, T.CONTENT_TOP, Inches(0.055), T.CONTENT_H,
               fill=T.status_color(status))
        T.add_textbox(slide, x + Inches(0.28), T.CONTENT_TOP + Inches(0.20),
                      col_w - Inches(0.56), Inches(0.28),
                      heading, size=Pt(13), color=T.status_color(status), bold=True)
        T.bullet_list(slide, x + Inches(0.28), T.CONTENT_TOP + Inches(0.62),
                      col_w - Inches(0.56), T.CONTENT_H - Inches(0.85),
                      items[:9] or [empty_text], size=Pt(11))

    n_high = sum(1 for i in eq_infos.values() if i['status'] == 'High')
    n_alarm = sum(1 for i in eq_infos.values() if i['status'] == 'Alarm')
    _column(T.MARGIN_X, f"Prioritas Tinggi - {n_high} equipment", 'High',
            buckets['High'], "Tidak ada equipment berstatus High pada periode ini.")
    _column(T.MARGIN_X + col_w + col_gap, f"Perhatian - {n_alarm} equipment", 'Alarm',
            buckets['Alarm'], "Tidak ada equipment berstatus Alarm pada periode ini.")
    return references


def _build_method_slide(deck, references, meta_right):
    slide = deck.slide("Metodologi dan Catatan Pengambilan Data", meta_right=meta_right)
    col_gap = Inches(0.25)
    col_w = int((T.CONTENT_W - col_gap) / 2)

    T.rounded_panel(slide, T.MARGIN_X, T.CONTENT_TOP, col_w, T.CONTENT_H,
                    fill=T.PANEL, line=T.BORDER)
    T.add_textbox(slide, T.MARGIN_X + Inches(0.28), T.CONTENT_TOP + Inches(0.20),
                  col_w - Inches(0.56), Inches(0.28),
                  "Prosedur Pengukuran (ATPOLL II)", size=Pt(13), color=T.BRAND, bold=True)
    T.bullet_list(slide, T.MARGIN_X + Inches(0.28), T.CONTENT_TOP + Inches(0.62),
                  col_w - Inches(0.56), T.CONTENT_H - Inches(0.85),
                  METODOLOGI_LINES, size=Pt(11))

    x2 = T.MARGIN_X + col_w + col_gap
    T.rounded_panel(slide, x2, T.CONTENT_TOP, col_w, T.CONTENT_H,
                    fill=T.PANEL, line=T.BORDER)
    T.add_textbox(slide, x2 + Inches(0.28), T.CONTENT_TOP + Inches(0.20),
                  col_w - Inches(0.56), Inches(0.28),
                  "Kerangka Tindak Lanjut", size=Pt(13), color=T.BRAND, bold=True)
    items = list(TINDAK_LANJUT_LINES)
    if references:
        items.append("Referensi: " + "; ".join(references[:3]))
    T.bullet_list(slide, x2 + Inches(0.28), T.CONTENT_TOP + Inches(0.62),
                  col_w - Inches(0.56), T.CONTENT_H - Inches(0.85),
                  items, size=Pt(11))
    return slide


# --------------------------------------------------------------------------
# Entry point
# --------------------------------------------------------------------------

def create_ppt(df_latest, context=None, history_df=None, include_images=True,
               include_trend=True, progress=None):
    """Bangun laporan PPTX dari data MCSA terbaru.

    df_latest      : DataFrame long-format berisi nilai terakhir per equipment.
    context        : dict opsional (period_start, period_end, unit_label, volt_label).
    history_df     : DataFrame riwayat untuk grafik tren (opsional).
    include_images : sisipkan gambar spektrum pada equipment Alarm/High.
    include_trend  : sisipkan slide tren parameter.
    progress       : callable(done, total, label) opsional untuk progress bar.

    Return io.BytesIO berisi file .pptx.
    """
    context = context or {}
    period_start = context.get('period_start')
    period_end = context.get('period_end')
    unit_label = context.get('unit_label')
    volt_label = context.get('volt_label')

    period_text = f"Periode: {period_start} s/d {period_end}" if (period_start and period_end) else None
    scope_text = " | ".join([p for p in [unit_label, volt_label] if p]) or None
    footer_bits = [b for b in [period_text, scope_text] if b]
    footer_text = "   |   ".join(footer_bits)
    meta_right = scope_text or "MCSA Assistant"

    deck = _Deck(footer_text=footer_text)

    subtitle_lines = []
    if period_text:
        subtitle_lines.append(period_text)
    if scope_text:
        subtitle_lines.append(scope_text)
    subtitle_lines.append("Motor Current Signature Analysis - ATPOLL II")
    deck.cover(
        "Laporan MCSA\n& Kondisi Equipment",
        subtitle_lines,
        f"Dibuat oleh MCSA Assistant  ·  {datetime.now().strftime('%d-%m-%Y %H:%M')}",
    )

    if not isinstance(df_latest, pd.DataFrame) or df_latest.empty or 'Parameter' not in df_latest.columns:
        slide = deck.slide("Ringkasan Eksekutif", meta_right=meta_right)
        T.add_textbox(slide, T.MARGIN_X, T.CONTENT_TOP, T.CONTENT_W, Inches(0.5),
                      "Tidak ada data pada filter dan periode yang dipilih.",
                      size=Pt(13), color=T.MUTED)
        return deck.save()

    thresholds = _thresholds()
    mapping = _load_equipment_mapping()
    img_index = None
    if include_images:
        try:
            img_index = SpectrumImageIndex()
        except Exception:
            img_index = None

    # --- Metadata & status per equipment ---------------------------------
    eq_infos = {}
    for eq, grp in df_latest.groupby('Equipment', sort=False):
        cond = grp[grp['Parameter'].astype(str) == 'Kondisi']
        raw_status = None
        if not cond.empty:
            for col in ('Status_Category', 'Status', 'Raw_Value'):
                if col in cond.columns and pd.notna(cond[col].iloc[0]):
                    raw_status = cond[col].iloc[0]
                    break
        status = T.canon_status(raw_status)

        full_name = _clean_text(grp['Full_Name'].iloc[0]) if 'Full_Name' in grp.columns else ''
        display = full_name or str(eq)
        long_name = _equipment_long_name(eq, mapping)
        # Full_Name kadang sudah deskriptif ("Closed Cooling Water Pump 1B").
        # Menambahkan nama panjang di situ hanya menghasilkan judul ganda,
        # jadi lewati bila Full_Name sudah berisi >=2 kata alfabetik.
        descriptive = len([w for w in display.split() if w.isalpha() and len(w) > 2]) >= 2
        if long_name and (descriptive or long_name.lower() in display.lower()):
            long_name = None
        title = f"{display} - {long_name}" if long_name else display
        title = _shorten(title, 66)

        dates = pd.to_datetime(grp.get('Date'), errors='coerce').dropna() if 'Date' in grp.columns else pd.Series(dtype='datetime64[ns]')
        last_date = dates.max() if not dates.empty else None

        eq_infos[eq] = {
            'title': title,
            'title_short': display,
            'status': status,
            'unit': _clean_text(grp['Unit_Name'].iloc[0]) if 'Unit_Name' in grp.columns else '',
            'volt': _clean_text(grp['Voltage_Level'].iloc[0]) if 'Voltage_Level' in grp.columns else '',
            'date': last_date.strftime('%Y-%m-%d') if last_date is not None else None,
            'date_txt': f"Pengukuran {last_date.strftime('%d %b %Y')}" if last_date is not None else '',
        }

    status_counts = {}
    for info in eq_infos.values():
        status_counts[info['status']] = status_counts.get(info['status'], 0) + 1
    total_eq = len(eq_infos)

    risk_rows = []
    try:
        risk_rows = build_risk_summary(df_latest, top_n=10) or []
    except Exception:
        risk_rows = []

    _build_summary_slide(deck, status_counts, total_eq, risk_rows, meta_right,
                         eq_infos=eq_infos)
    _build_risk_slide(deck, risk_rows, meta_right)

    cond_df = pd.DataFrame([
        {'Equipment': eq, '_status': i['status'], '_unit': i['unit'] or 'Lainnya',
         '_volt': i['volt'] or '-'}
        for eq, i in eq_infos.items()
    ])
    _build_distribution_slide(deck, cond_df, meta_right)

    # --- Detail per equipment, dikelompokkan per Unit/Voltage -------------
    def _group_key(eq):
        i = eq_infos[eq]
        return (i['unit'] or 'zzz', i['volt'] or 'zzz')

    status_rank = {'High': 0, 'Alarm': 1, 'Monitoring': 2, 'Normal': 3, 'Standby': 4, 'Unknown': 5}
    equipments = sorted(eq_infos.keys(),
                        key=lambda e: (_group_key(e), status_rank.get(eq_infos[e]['status'], 9), str(e)))

    total_steps = len(equipments)
    current_group = None
    for done, eq in enumerate(equipments, start=1):
        info = eq_infos[eq]
        key = _group_key(eq)
        if key != current_group:
            current_group = key
            members = [e for e in equipments if _group_key(e) == key]
            counts = {}
            for m in members:
                s = eq_infos[m]['status']
                counts[s] = counts.get(s, 0) + 1
            summary = "   ".join(f"{s}: {counts[s]}" for s in T.STATUS_ORDER if counts.get(s))
            unit_txt = info['unit'] or 'Lainnya'
            volt_txt = info['volt'] or '-'
            deck.divider(unit_txt, volt_txt,
                         [f"{len(members)} equipment", summary])

        eq_data = df_latest[df_latest['Equipment'] == eq].copy()
        eq_data = eq_data.dropna(subset=['Parameter'])
        eq_data['Parameter'] = eq_data['Parameter'].astype(str)
        eq_data = eq_data[~eq_data['Parameter'].str.startswith('Ringkasan Kinerja')]
        order_map = {p: i for i, p in enumerate(PREFERRED_ORDER)}
        eq_data['_k'] = eq_data['Parameter'].map(lambda p: order_map.get(p, 10_000))
        eq_data = eq_data.sort_values(by=['_k', 'Parameter']).drop(columns=['_k'])
        eq_data = eq_data.drop_duplicates(subset=['Parameter'], keep='first')

        _build_detail_slides(deck, eq, eq_data, info, thresholds, img_index,
                             include_images, meta_right)

        if include_trend and info['status'] in {'Alarm', 'High'}:
            _build_trend_slide(deck, eq, info, history_df, thresholds, meta_right)

        if callable(progress):
            try:
                progress(done, total_steps, info['title_short'])
            except Exception:
                pass

    references = _build_recommendation_slides(deck, df_latest, eq_infos, meta_right)
    _build_method_slide(deck, references, meta_right)

    return deck.save()


if __name__ == "__main__":
    from src.data_loader import load_mcsa_data, get_latest_data, get_data_path

    df = load_mcsa_data(get_data_path('mcsa_updated.csv'))
    latest = get_latest_data(df)
    ppt_io = create_ppt(latest, history_df=df)
    with open("test_output.pptx", "wb") as f:
        f.write(ppt_io.read())
    print("PPT Generated successfully.")
