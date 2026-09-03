"""Cross-domain engineering condition report page."""

from datetime import date
from io import BytesIO

import pandas as pd
from docx import Document

from src.asset_registry import MONITORING_MODULES, list_assets, load_condition_history
from src.components.theme import render_page_header


def _build_docx(records: pd.DataFrame, start: date, end: date, module: str) -> BytesIO:
    document = Document()
    document.add_heading("Laporan kondisi engineering", level=0)
    document.add_paragraph(f"Periode: {start:%d-%m-%Y} s/d {end:%d-%m-%Y}")
    document.add_paragraph(f"Modul: {module}")
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, ("Tanggal", "Aset", "Modul", "Kondisi", "Ringkasan", "Bukti")):
        cell.text = label
    for _, row in records.sort_values("test_date").iterrows():
        cells = table.add_row().cells
        values = (row["test_date"].strftime("%d-%m-%Y"), row["asset_name"], row["module"], row["condition"], row["summary"], row["source_file"])
        for cell, value in zip(cells, values):
            cell.text = str(value or "-")
    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output


def render_asset_reports_page(st) -> None:
    render_page_header(st, "Laporan kondisi", "Laporan DGA, MCSA, Vibrasi, PD, Tribology, dan Thermal berdasarkan periode.")
    history = load_condition_history()
    assets = list_assets()
    if history.empty:
        st.info("Belum ada riwayat kondisi dari Asset Management. Tambahkan hasil pengujian terlebih dahulu.")
        return

    min_date = history["test_date"].min().date()
    max_date = history["test_date"].max().date()
    with st.form("asset_report_filter", border=True):
        c1, c2, c3 = st.columns(3)
        with c1:
            selected_module = st.selectbox("Metode pengujian", ["Semua"] + list(MONITORING_MODULES))
        with c2:
            selected_assets = st.multiselect("Aset", [asset["asset_id"] for asset in assets], format_func=lambda asset_id: next((f"{a['name']} ({asset_id})" for a in assets if a["asset_id"] == asset_id), asset_id))
        with c3:
            selected_dates = st.date_input("Periode laporan", value=(min_date, max_date), min_value=min_date, max_value=max_date)
        applied = st.form_submit_button("Terapkan filter", icon=":material/filter_alt:")

    if not applied:
        selected_module, selected_assets, selected_dates = "Semua", [], (min_date, max_date)
    start, end = selected_dates if isinstance(selected_dates, tuple) else (selected_dates, selected_dates)
    records = history[(history["test_date"].dt.date >= start) & (history["test_date"].dt.date <= end)].copy()
    if selected_module != "Semua":
        records = records[records["module"] == selected_module]
    if selected_assets:
        records = records[records["asset_id"].isin(selected_assets)]
    st.metric("Hasil pengujian dalam laporan", len(records))
    if records.empty:
        st.warning("Tidak ada hasil pengujian pada filter tersebut.")
        return
    st.dataframe(records.drop(columns=["record_id", "created_at"]), hide_index=True, width="stretch")
    st.download_button("Unduh data CSV", records.to_csv(index=False).encode("utf-8-sig"), f"Laporan_Kondisi_{start:%Y%m%d}_{end:%Y%m%d}.csv", "text/csv", icon=":material/download:")
    document = _build_docx(records, start, end, selected_module)
    st.download_button("Unduh laporan Word", document, f"Laporan_Kondisi_{start:%Y%m%d}_{end:%Y%m%d}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", icon=":material/description:")
