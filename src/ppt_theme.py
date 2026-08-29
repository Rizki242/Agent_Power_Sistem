"""Tema visual untuk laporan PowerPoint MCSA.

Modul ini memusatkan palet warna, font, geometri, dan primitive gambar agar
seluruh slide punya tampilan konsisten. Tidak ada file template .pptx eksternal:
kanvas 16:9 dan seluruh dekorasi dibangun secara programatis.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.chart import XL_LEGEND_POSITION
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement

# --- Kanvas ---------------------------------------------------------------
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

MARGIN_X = Inches(0.55)
HEADER_H = Inches(1.05)          # tinggi blok judul
CONTENT_TOP = Inches(1.30)       # batas atas area konten
FOOTER_TOP = Inches(6.95)
CONTENT_BOTTOM = FOOTER_TOP - Inches(0.12)
CONTENT_W = SLIDE_W - 2 * MARGIN_X
CONTENT_H = CONTENT_BOTTOM - CONTENT_TOP

# --- Font -----------------------------------------------------------------
FONT = "Segoe UI"

# --- Palet ----------------------------------------------------------------
BRAND = RGBColor(0x0F, 0x4C, 0x81)
BRAND_DARK = RGBColor(0x0A, 0x33, 0x57)
ACCENT = RGBColor(0x0E, 0xA5, 0xE9)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)
SURFACE = RGBColor(0xFF, 0xFF, 0xFF)
PANEL = RGBColor(0xF8, 0xFA, 0xFC)
BORDER = RGBColor(0xE2, 0xE8, 0xF0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# Satu sumber kebenaran warna status, dipakai PPT maupun DOCX.
STATUS_HEX = {
    'Normal': 0x16A34A,
    'Monitoring': 0x0EA5E9,
    'Alarm': 0xF59E0B,
    'High': 0xDC2626,
    'Standby': 0x94A3B8,
    'Unknown': 0xCBD5E1,
}
STATUS_TINT_HEX = {
    'Normal': 0xDCFCE7,
    'Monitoring': 0xE0F2FE,
    'Alarm': 0xFEF3C7,
    'High': 0xFEE2E2,
    'Standby': 0xF1F5F9,
    'Unknown': 0xF1F5F9,
}

STATUS_ORDER = ['Normal', 'Alarm', 'High', 'Standby', 'Unknown']


def status_color(status: str) -> RGBColor:
    return RGBColor.from_string(f"{STATUS_HEX.get(status, STATUS_HEX['Unknown']):06X}")


def status_tint(status: str) -> RGBColor:
    return RGBColor.from_string(f"{STATUS_TINT_HEX.get(status, STATUS_TINT_HEX['Unknown']):06X}")


def canon_status(value) -> str:
    """Kanonikalisasi status mentah ke salah satu STATUS_ORDER."""
    v = str(value if value is not None else '').strip().lower()
    if v in {'', 'nan', 'none', 'unknown'}:
        return 'Unknown'
    if v == 'normal':
        return 'Normal'
    if v in {'alarm', 'warning'}:
        return 'Alarm'
    if v in {'high', 'critical'}:
        return 'High'
    if v == 'standby':
        return 'Standby'
    if 'standby' in v:
        return 'Standby'
    if 'high' in v or 'critical' in v or 'trip' in v or 'bad' in v:
        return 'High'
    if 'alarm' in v or 'warning' in v:
        return 'Alarm'
    if 'normal' in v:
        return 'Normal'
    return 'Unknown'


# --- Primitive teks -------------------------------------------------------

def set_text(text_frame, lines, size=Pt(12), color=INK, bold=False,
             align=PP_ALIGN.LEFT, space_after=Pt(0), line_spacing=None):
    """Isi text_frame dengan daftar baris, mengeset font per-run.

    python-pptx tidak mewariskan font master ke textbox yang dibuat manual,
    jadi setiap run harus diberi nama font secara eksplisit.
    `lines` boleh berupa str atau list[str] atau list[dict] dengan kunci
    text/size/color/bold/align.
    """
    if isinstance(lines, str):
        lines = [lines]
    lines = list(lines) or ['']

    text_frame.word_wrap = True
    for idx, item in enumerate(lines):
        spec = item if isinstance(item, dict) else {'text': item}
        para = text_frame.paragraphs[0] if idx == 0 else text_frame.add_paragraph()
        para.alignment = spec.get('align', align)
        para.space_after = spec.get('space_after', space_after)
        if line_spacing is not None:
            para.line_spacing = line_spacing
        run = para.add_run()
        run.text = str(spec.get('text', ''))
        font = run.font
        font.name = FONT
        font.size = spec.get('size', size)
        font.bold = spec.get('bold', bold)
        font.color.rgb = spec.get('color', color)
    return text_frame


def add_textbox(slide, left, top, width, height, lines, **kwargs):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    set_text(tf, lines, **kwargs)
    return box


# --- Primitive bentuk -----------------------------------------------------

def _send_to_back(slide, shape):
    """Pindahkan shape ke urutan paling belakang pada spTree."""
    el = shape._element
    tree = slide.shapes._spTree
    tree.remove(el)
    # index 0..1 dipakai nvGrpSpPr dan grpSpPr; shape pertama mulai di index 2
    tree.insert(2, el)


def fill_background(slide, color=SURFACE):
    """Gambar rect full-bleed lalu kirim ke belakang.

    `slide.background.fill` tidak konsisten diterapkan per-slide oleh
    PowerPoint saat layout Blank dipakai, jadi rect eksplisit lebih andal.
    """
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    _send_to_back(slide, shape)
    return shape


def rect(slide, left, top, width, height, fill=None, line=None, line_w=Pt(1)):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = line_w
    shape.shadow.inherit = False
    shape.text_frame.word_wrap = True
    return shape


def rounded_panel(slide, left, top, width, height, fill=PANEL, line=BORDER, radius=0.06):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    try:
        shape.adjustments[0] = radius
    except (IndexError, ValueError):
        pass
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    shape.shadow.inherit = False
    shape.text_frame.word_wrap = True
    return shape


def status_chip(slide, left, top, text, status, width=Inches(1.35), height=Inches(0.34)):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    try:
        shape.adjustments[0] = 0.5
    except (IndexError, ValueError):
        pass
    shape.fill.solid()
    shape.fill.fore_color.rgb = status_color(status)
    shape.line.fill.background()
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_text(tf, str(text).upper(), size=Pt(11), color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    return shape


def kpi_card(slide, left, top, width, height, value, label, status):
    panel = rounded_panel(slide, left, top, width, height,
                          fill=status_tint(status), line=None)
    accent = rect(slide, left, top, Inches(0.055), height, fill=status_color(status))
    add_textbox(slide, left + Inches(0.22), top + Inches(0.14),
                width - Inches(0.34), Inches(0.62),
                str(value), size=Pt(34), color=status_color(status), bold=True)
    add_textbox(slide, left + Inches(0.22), top + Inches(0.76),
                width - Inches(0.34), Inches(0.28),
                str(label).upper(), size=Pt(10), color=MUTED, bold=True)
    return panel, accent


def bullet_list(slide, left, top, width, height, items, size=Pt(12),
                color=INK, bullet_color=ACCENT, space_after=Pt(7)):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    items = list(items) or ['-']
    for idx, item in enumerate(items):
        para = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        para.space_after = space_after
        dot = para.add_run()
        dot.text = "▪  "
        dot.font.name = FONT
        dot.font.size = size
        dot.font.bold = True
        dot.font.color.rgb = bullet_color
        run = para.add_run()
        run.text = str(item)
        run.font.name = FONT
        run.font.size = size
        run.font.color.rgb = color
    return box


# --- Tabel ----------------------------------------------------------------

def _clear_table_style(table):
    """Matikan banding/emphasis bawaan agar fill manual terlihat."""
    tbl_pr = table._tbl.tblPr
    for attr in ('firstRow', 'lastRow', 'firstCol', 'lastCol', 'bandRow', 'bandCol'):
        tbl_pr.set(attr, '0')
    # Hapus referensi tema tabel supaya tidak ada style biru bawaan.
    for style_id in tbl_pr.findall(qn('a:tableStyleId')):
        tbl_pr.remove(style_id)


_BORDER_TAGS = ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB')


def _set_cell_border(cell, color, width_pt=0.75):
    """Ganti garis sel bawaan (hitam) dengan warna tema.

    python-pptx tidak punya API border tabel, jadi elemen a:lnL/R/T/B ditulis
    manual. Skema OOXML mengharuskan elemen ini menjadi anak PERTAMA dari
    a:tcPr (mendahului a:solidFill), sehingga disisipkan di indeks awal.
    """
    tc_pr = cell._tc.get_or_add_tcPr()
    for tag in _BORDER_TAGS:
        existing = tc_pr.find(qn(tag))
        if existing is not None:
            tc_pr.remove(existing)
    for i, tag in enumerate(_BORDER_TAGS):
        ln = OxmlElement(tag)
        ln.set('w', str(int(width_pt * 12700)))
        ln.set('cap', 'flat')
        ln.set('cmpd', 'sng')
        ln.set('algn', 'ctr')
        solid = OxmlElement('a:solidFill')
        srgb = OxmlElement('a:srgbClr')
        srgb.set('val', str(color))
        solid.append(srgb)
        ln.append(solid)
        tc_pr.insert(i, ln)


def style_table(table, col_widths=None, font_size=Pt(10), row_h=Inches(0.26),
                header_fill=BRAND, header_color=WHITE, align=None,
                highlight_rows=None, header_size=None, border_color=BORDER):
    """Terapkan tema pada tabel.

    col_widths   : list lebar kolom (Emu/Inches)
    align        : list PP_ALIGN per kolom
    highlight_rows: dict {row_idx (1-based data row): status} untuk mewarnai baris
    """
    _clear_table_style(table)
    highlight_rows = highlight_rows or {}
    n_cols = len(table.columns)

    if col_widths:
        for i, w in enumerate(col_widths[:n_cols]):
            table.columns[i].width = Emu(int(w))

    for r_idx, row in enumerate(table.rows):
        row.height = Emu(int(row_h))
        is_header = (r_idx == 0)
        hl = highlight_rows.get(r_idx)
        for c_idx, cell in enumerate(row.cells):
            cell.margin_left = Inches(0.07)
            cell.margin_right = Inches(0.07)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE

            _set_cell_border(cell, header_fill if is_header else border_color)

            cell.fill.solid()
            if is_header:
                cell.fill.fore_color.rgb = header_fill
            elif hl:
                cell.fill.fore_color.rgb = status_tint(hl)
            else:
                cell.fill.fore_color.rgb = SURFACE if (r_idx % 2 == 1) else PANEL

            tf = cell.text_frame
            tf.word_wrap = True
            for para in tf.paragraphs:
                if align and c_idx < len(align):
                    para.alignment = align[c_idx]
                else:
                    para.alignment = PP_ALIGN.LEFT
                if not para.runs:
                    continue
                for run in para.runs:
                    font = run.font
                    font.name = FONT
                    font.size = (header_size or font_size) if is_header else font_size
                    font.bold = bool(is_header)
                    if is_header:
                        font.color.rgb = header_color
                    elif hl:
                        font.color.rgb = status_color(hl)
                    else:
                        font.color.rgb = INK
    return table


def add_table(slide, rows, cols, left, top, width, height):
    return slide.shapes.add_table(rows, cols, left, top, width, height).table


# --- Chart ----------------------------------------------------------------

def style_chart(chart, legend=True, legend_pos=XL_LEGEND_POSITION.BOTTOM, font_size=Pt(9)):
    chart.font.name = FONT
    chart.font.size = font_size
    chart.font.color.rgb = INK
    chart.has_title = False
    chart.has_legend = bool(legend)
    if legend:
        chart.legend.position = legend_pos
        chart.legend.include_in_layout = False
    return chart


# --- Slide ----------------------------------------------------------------

def new_presentation():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _blank_layout(prs):
    # Layout index 6 pada template default python-pptx adalah "Blank".
    return prs.slide_layouts[6]


def add_slide(prs, title=None, meta_right=None, footer_left=None, background=SURFACE):
    """Slide standar: background, accent bar, judul, meta kanan, footer."""
    slide = prs.slides.add_slide(_blank_layout(prs))
    fill_background(slide, background)

    if title:
        title = str(title)
        # Auto-shrink supaya judul panjang tetap satu baris dan tidak menimpa
        # subjudul/garis pemisah di bawahnya.
        if len(title) <= 46:
            t_size = Pt(24)
        elif len(title) <= 62:
            t_size = Pt(20)
        else:
            t_size = Pt(17)
        rect(slide, MARGIN_X, Inches(0.46), Inches(0.075), Inches(0.44), fill=ACCENT)
        box = add_textbox(slide, MARGIN_X + Inches(0.20), Inches(0.42),
                          CONTENT_W - Inches(3.6), Inches(0.52),
                          title, size=t_size, color=INK, bold=True)
        box.text_frame.word_wrap = False

    if meta_right:
        add_textbox(slide, SLIDE_W - MARGIN_X - Inches(3.4), Inches(0.52),
                    Inches(3.4), Inches(0.40),
                    str(meta_right), size=Pt(10.5), color=MUTED,
                    align=PP_ALIGN.RIGHT)

    if title or meta_right:
        rect(slide, MARGIN_X, HEADER_H + Inches(0.06), CONTENT_W, Pt(1), fill=BORDER)

    if footer_left:
        add_textbox(slide, MARGIN_X, FOOTER_TOP, CONTENT_W - Inches(1.2), Inches(0.26),
                    str(footer_left), size=Pt(9), color=MUTED)

    return slide


def add_page_number(slide, number):
    add_textbox(slide, SLIDE_W - MARGIN_X - Inches(1.0), FOOTER_TOP,
                Inches(1.0), Inches(0.26), str(number),
                size=Pt(9), color=MUTED, align=PP_ALIGN.RIGHT)


def add_divider_slide(prs, title, subtitle=None, lines=None):
    """Slide pemisah seksi: panel brand full-bleed."""
    slide = prs.slides.add_slide(_blank_layout(prs))
    fill_background(slide, BRAND)
    rect(slide, 0, SLIDE_H - Inches(0.30), SLIDE_W, Inches(0.30), fill=ACCENT)

    add_textbox(slide, MARGIN_X + Inches(0.6), Inches(2.75), CONTENT_W - Inches(1.2),
                Inches(0.85), str(title), size=Pt(38), color=WHITE, bold=True)
    if subtitle:
        add_textbox(slide, MARGIN_X + Inches(0.6), Inches(3.65), CONTENT_W - Inches(1.2),
                    Inches(0.40), str(subtitle), size=Pt(15),
                    color=RGBColor(0xBF, 0xDB, 0xFE))
    if lines:
        add_textbox(slide, MARGIN_X + Inches(0.6), Inches(4.20), CONTENT_W - Inches(1.2),
                    Inches(0.90), list(lines), size=Pt(12),
                    color=RGBColor(0xDB, 0xEA, 0xFE), space_after=Pt(4))
    return slide


def add_cover_slide(prs, title, subtitle_lines=None, footer=None):
    slide = prs.slides.add_slide(_blank_layout(prs))
    fill_background(slide, BRAND_DARK)

    # Blok aksen dekoratif di kanan.
    rect(slide, SLIDE_W - Inches(3.6), 0, Inches(3.6), SLIDE_H, fill=BRAND)
    rect(slide, SLIDE_W - Inches(3.6), 0, Inches(0.06), SLIDE_H, fill=ACCENT)

    rect(slide, MARGIN_X + Inches(0.4), Inches(2.30), Inches(1.5), Inches(0.075), fill=ACCENT)
    add_textbox(slide, MARGIN_X + Inches(0.4), Inches(2.60), Inches(8.4), Inches(1.5),
                str(title), size=Pt(40), color=WHITE, bold=True, line_spacing=1.05)
    if subtitle_lines:
        add_textbox(slide, MARGIN_X + Inches(0.4), Inches(4.25), Inches(8.4), Inches(1.4),
                    list(subtitle_lines), size=Pt(14),
                    color=RGBColor(0xBF, 0xDB, 0xFE), space_after=Pt(6))
    if footer:
        add_textbox(slide, MARGIN_X + Inches(0.4), Inches(6.55), Inches(8.4), Inches(0.3),
                    str(footer), size=Pt(10), color=RGBColor(0x93, 0xC5, 0xFD))
    return slide
